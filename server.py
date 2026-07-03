"""AutoLAB web server.

Run as a web app:
    python -m uvicorn server:app --reload                       (localhost only)
    python -m uvicorn server:app --host 0.0.0.0 --port 8000     (LAN/phone access)

Run as a desktop app:
    python desktop.py

Phone access on the same Wi-Fi: find your laptop's LAN IP (e.g. via
`ipconfig`), then visit http://<that-ip>:8000 from the phone's browser.

The same backend powers both. The desktop launcher sets AUTOLAB_DESKTOP=1 so
that the /api/profile and other local-only endpoints are unlocked; the web
mode returns 403 from those routes.
"""

from __future__ import annotations

import io
import json
import os
import sys
import tempfile
import urllib.error
import urllib.request
from pathlib import Path

from docx import Document
from fastapi import FastAPI, File, Form, Header, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

import billing
from autolab import (
    apply_block_edits_bytes,
    extract_blocks,
    replace_in_doc,
    replace_in_docx_bytes,
)
from assemble import add_solved_question_bytes
from gemini import gemini_call, patch_docx_images, strip_fences
from labshot import render_code_shot, render_terminal_shot
from runner import LANGS, run_pair, run_single


def _resource_root() -> Path:
    """Where bundled resources (e.g. static/) live.

    PyInstaller extracts data files to sys._MEIPASS at runtime; in dev we use
    the source directory.
    """
    if getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"):
        return Path(sys._MEIPASS)  # type: ignore[attr-defined]
    return Path(__file__).resolve().parent


ROOT = _resource_root()
STATIC = ROOT / "static"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

DESKTOP_MODE = os.environ.get("AUTOLAB_DESKTOP") == "1"
DATA_DIR = Path(os.environ.get("APPDATA") or Path.home() / ".config") / "AutoLAB"
PROFILE_PATH = DATA_DIR / "profile.json"
GENERATED_DIR = DATA_DIR / "Generated"

if DESKTOP_MODE:
    try:
        GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    except Exception:
        pass

DEFAULT_PROFILE = {
    "saved_urks": [],
    "default_urk": "",
    "ai_api_key": "",
    "preferences": {
        "default_course_title": "",
        "default_template": "",
    },
}


def _public_profile(p: dict) -> dict:
    """Strip the raw API key before returning to the client."""
    safe = {k: v for k, v in p.items() if k != "ai_api_key"}
    raw = p.get("ai_api_key", "")
    safe["ai_key_set"] = bool(raw)
    safe["ai_key_preview"] = ("…" + raw[-4:]) if raw else ""
    return safe


def _load_profile() -> dict:
    if not PROFILE_PATH.exists():
        return dict(DEFAULT_PROFILE)
    try:
        data = json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
    except Exception:
        return dict(DEFAULT_PROFILE)
    merged = dict(DEFAULT_PROFILE)
    merged.update(data)
    # Ensure nested defaults exist
    merged["preferences"] = {**DEFAULT_PROFILE["preferences"], **merged.get("preferences", {})}
    return merged


def _save_profile(profile: dict) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    PROFILE_PATH.write_text(json.dumps(profile, indent=2), encoding="utf-8")


def _require_desktop() -> None:
    if not DESKTOP_MODE:
        raise HTTPException(403, "This feature is available only in the desktop app.")


app = FastAPI(title="AutoLAB", docs_url=None, redoc_url=None)


@app.get("/")
async def index() -> FileResponse:
    return FileResponse(STATIC / "index.html")


# Legal pages — Razorpay activation requires these to exist on the site.
def _legal_route(page: str):
    async def _serve() -> FileResponse:
        return FileResponse(STATIC / "legal" / f"{page}.html")
    return _serve


for _page in ("terms", "privacy", "refunds", "contact"):
    app.get(f"/{_page}", include_in_schema=False)(_legal_route(_page))


@app.post("/api/replace")
async def replace(
    file: UploadFile = File(...),
    find: str = Form(...),
    replace: str = Form(""),
    patch_images: bool = Form(False),
    api_key: str = Form(""),
    authorization: str = Header(""),
) -> Response:
    if not file.filename or not file.filename.lower().endswith(".docx"):
        raise HTTPException(400, "Please upload a .docx file.")
    if not find:
        raise HTTPException(400, "Find text is required.")

    data = await file.read()
    if not data:
        raise HTTPException(400, "Uploaded file is empty.")

    try:
        out_bytes, count = replace_in_docx_bytes(data, find, replace)
    except Exception as e:
        raise HTTPException(500, f"Failed to process document: {e}") from e

    images_patched = images_scanned = 0
    if patch_images:
        key, charged = _paid_key(authorization, api_key, "patch_screens")
        try:
            out_bytes, images_patched, images_scanned = patch_docx_images(
                out_bytes, find, replace, key
            )
        except urllib.error.HTTPError as e:
            _refund(charged, "patch_screens")
            msg = e.read().decode("utf-8", errors="replace") if e.fp else str(e)
            raise HTTPException(502, f"Gemini API error {e.code}: {msg[:200]}") from e
        except Exception as e:
            _refund(charged, "patch_screens")
            raise HTTPException(502, f"Image patching failed: {e}") from e

    base = Path(file.filename).stem or "document"
    out_name = f"{base}_modified.docx"
    return Response(
        content=out_bytes,
        media_type=DOCX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{out_name}"',
            "X-Replace-Count": str(count),
            "X-Images-Patched": str(images_patched),
            "X-Images-Scanned": str(images_scanned),
            "Access-Control-Expose-Headers": (
                "X-Replace-Count, X-Images-Patched, X-Images-Scanned, Content-Disposition"
            ),
        },
    )


def _profile_key() -> str:
    """Gemini key from the desktop profile; empty string in web mode."""
    return _load_profile().get("ai_api_key", "") if DESKTOP_MODE else ""


# --- billing / credits --------------------------------------------------

def _require_billing() -> None:
    if not billing.enabled():
        raise HTTPException(503, "Payments are not configured on this server.")


def _auth(authorization: str) -> dict:
    try:
        return billing.auth_user(authorization)
    except ValueError as e:
        raise HTTPException(401, str(e)) from e


def _paid_key(authorization: str, api_key: str, op: str) -> tuple[str, str]:
    """Resolve which Gemini key a paid op runs on.

    BYO key (form field, or the desktop profile) -> free, no account needed.
    Otherwise the op needs a signed-in user with enough credits: we spend
    upfront and the caller refunds via _refund() if the op then fails.
    Returns (gemini_key, charged_user_id) — user_id is "" when free.
    """
    key = api_key.strip() or _profile_key()
    if key:
        return key, ""
    if not billing.enabled():
        raise HTTPException(
            400,
            "This needs a Gemini API key — free at aistudio.google.com.",
        )
    user = _auth(authorization)
    if not billing.SERVER_GEMINI_KEY:
        raise HTTPException(503, "Server AI key is not configured.")
    price = billing.PRICES[op]
    try:
        left = billing.spend(user["id"], price, op)
    except Exception as e:
        raise HTTPException(502, f"Credit check failed: {e}") from e
    if left < 0:
        raise HTTPException(
            402, f"Not enough credits — this costs ₹{price}. Top up in Account."
        )
    return billing.SERVER_GEMINI_KEY, user["id"]


def _refund(user_id: str, op: str) -> None:
    if not user_id:
        return
    try:
        billing.credit(user_id, billing.PRICES[op], f"refund:{op}")
    except Exception:
        # ponytail: a lost refund shows as a gap in the ledger — handle by hand
        # if a user ever reports one.
        pass


@app.get("/api/billing/config")
async def billing_config() -> dict:
    """Public billing facts the frontend needs. No secrets here."""
    return {
        "enabled": billing.enabled(),
        "supabase_url": billing.SUPABASE_URL,
        "supabase_anon_key": billing.SUPABASE_ANON_KEY,
        "razorpay_key_id": billing.RAZORPAY_KEY_ID,
        "prices": billing.PRICES,
        "packs": billing.PACKS,
        "min_topup": billing.MIN_TOPUP,
        "max_topup": billing.MAX_TOPUP,
    }


@app.get("/api/me")
async def me(authorization: str = Header("")) -> dict:
    _require_billing()
    user = _auth(authorization)
    try:
        bal = billing.balance(user["id"])
    except Exception as e:
        raise HTTPException(502, f"Could not read balance: {e}") from e
    return {"email": user["email"], "balance": bal}


class OrderIn(BaseModel):
    amount: int


@app.post("/api/pay/order")
async def pay_order(body: OrderIn, authorization: str = Header("")) -> dict:
    _require_billing()
    user = _auth(authorization)
    if not (billing.MIN_TOPUP <= body.amount <= billing.MAX_TOPUP):
        raise HTTPException(
            400, f"Top-up must be ₹{billing.MIN_TOPUP}–₹{billing.MAX_TOPUP}."
        )
    try:
        order = billing.rzp_create_order(body.amount, user["id"])
    except Exception as e:
        raise HTTPException(502, f"Could not create payment order: {e}") from e
    return {
        "order_id": order["id"],
        "amount": order["amount"],
        "currency": "INR",
        "key_id": billing.RAZORPAY_KEY_ID,
        "email": user["email"],
    }


class VerifyIn(BaseModel):
    order_id: str
    payment_id: str
    signature: str


@app.post("/api/pay/verify")
async def pay_verify(body: VerifyIn, authorization: str = Header("")) -> dict:
    """Instant crediting right after checkout; the webhook is the backup."""
    _require_billing()
    user = _auth(authorization)
    if not billing.verify_payment_sig(body.order_id, body.payment_id, body.signature):
        raise HTTPException(400, "Payment signature check failed.")
    try:
        order = billing.rzp_get_order(body.order_id)
    except Exception as e:
        raise HTTPException(502, f"Could not fetch the order: {e}") from e
    if (order.get("notes") or {}).get("user_id") != user["id"]:
        raise HTTPException(403, "This order belongs to a different account.")
    rupees = int(order["amount"]) // 100
    try:
        bal = billing.credit(user["id"], rupees, "topup", ref=body.payment_id)
    except Exception as e:
        raise HTTPException(502, f"Payment ok but crediting failed: {e}") from e
    return {"balance": bal, "added": rupees}


@app.post("/api/pay/webhook")
async def pay_webhook(request: Request) -> dict:
    body = await request.body()
    sig = request.headers.get("X-Razorpay-Signature", "")
    if not billing.verify_webhook_sig(body, sig):
        raise HTTPException(400, "Bad webhook signature.")
    event = json.loads(body)
    if event.get("event") == "payment.captured":
        p = event["payload"]["payment"]["entity"]
        uid = (p.get("notes") or {}).get("user_id")
        if uid:
            billing.credit(uid, int(p["amount"]) // 100, "topup", ref=p["id"])
    return {"ok": True}


@app.get("/api/capabilities")
async def capabilities() -> dict:
    """Tell the frontend what the running backend can do."""
    return {
        "desktop": DESKTOP_MODE,
        "version": "0.8",
        "billing": billing.enabled(),
        # ponytail: no compilers/subprocess sandbox on Vercel — solve is local-only
        "solve": not os.environ.get("VERCEL"),
        "languages": list(LANGS),
        "data_dir": str(DATA_DIR) if DESKTOP_MODE else None,
        "generated_dir": str(GENERATED_DIR) if DESKTOP_MODE else None,
    }


class SaveUrk(BaseModel):
    urk: str


class ProfileUpdate(BaseModel):
    default_urk: str | None = None
    default_course_title: str | None = None
    ai_api_key: str | None = None


@app.get("/api/profile")
async def get_profile() -> dict:
    _require_desktop()
    return _public_profile(_load_profile())


@app.post("/api/profile/urks")
async def add_urk(body: SaveUrk) -> dict:
    _require_desktop()
    urk = body.urk.strip()
    if not urk:
        raise HTTPException(400, "URK can't be blank.")
    p = _load_profile()
    if urk not in p["saved_urks"]:
        p["saved_urks"].insert(0, urk)
        p["saved_urks"] = p["saved_urks"][:12]
    if not p["default_urk"]:
        p["default_urk"] = urk
    _save_profile(p)
    return _public_profile(p)


@app.delete("/api/profile/urks/{urk}")
async def remove_urk(urk: str) -> dict:
    _require_desktop()
    p = _load_profile()
    p["saved_urks"] = [u for u in p["saved_urks"] if u != urk]
    if p["default_urk"] == urk:
        p["default_urk"] = p["saved_urks"][0] if p["saved_urks"] else ""
    _save_profile(p)
    return _public_profile(p)


@app.patch("/api/profile")
async def update_profile(body: ProfileUpdate) -> dict:
    _require_desktop()
    p = _load_profile()
    if body.default_urk is not None:
        p["default_urk"] = body.default_urk.strip()
    if body.default_course_title is not None:
        p["preferences"]["default_course_title"] = body.default_course_title.strip()
    if body.ai_api_key is not None:
        p["ai_api_key"] = body.ai_api_key.strip()
    _save_profile(p)
    return _public_profile(p)


# --- Document editor endpoints ----------------------------------------------

class BlockEdit(BaseModel):
    id: str
    text: str


@app.post("/api/document/parse")
async def doc_parse(file: UploadFile = File(...)) -> dict:
    if not file.filename or not file.filename.lower().endswith(".docx"):
        raise HTTPException(400, "Please upload a .docx file.")
    data = await file.read()
    if not data:
        raise HTTPException(400, "Uploaded file is empty.")
    try:
        doc = Document(io.BytesIO(data))
        blocks = extract_blocks(doc)
    except Exception as e:
        raise HTTPException(500, f"Couldn't parse docx: {e}") from e
    return {
        "name": file.filename,
        "size": len(data),
        "block_count": len(blocks),
        "blocks": blocks,
    }


@app.post("/api/document/save")
async def doc_save(
    file: UploadFile = File(...),
    edits: str = Form(...),
) -> Response:
    try:
        parsed = json.loads(edits)
        if not isinstance(parsed, dict):
            raise ValueError("`edits` must be a JSON object {id: text}")
    except (ValueError, json.JSONDecodeError) as e:
        raise HTTPException(400, f"Invalid edits payload: {e}") from e

    data = await file.read()
    if not data:
        raise HTTPException(400, "Uploaded file is empty.")
    try:
        out_bytes, applied = apply_block_edits_bytes(data, parsed)
    except Exception as e:
        raise HTTPException(500, f"Failed to apply edits: {e}") from e

    base = Path(file.filename or "document").stem or "document"
    out_name = f"{base}_edited.docx"
    return Response(
        content=out_bytes,
        media_type=DOCX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{out_name}"',
            "X-Edits-Applied": str(applied),
            "Access-Control-Expose-Headers": "X-Edits-Applied, Content-Disposition",
        },
    )


@app.post("/api/document/add-question")
async def doc_add_question(
    file: UploadFile = File(...),
    number: str = Form(""),
    title: str = Form(...),
    description: str = Form(""),
) -> Response:
    """Append a new question (heading + description) to a docx."""
    if not file.filename or not file.filename.lower().endswith(".docx"):
        raise HTTPException(400, "Please upload a .docx file.")
    if not title.strip():
        raise HTTPException(400, "Title is required.")
    data = await file.read()
    if not data:
        raise HTTPException(400, "Uploaded file is empty.")
    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise HTTPException(500, f"Couldn't parse docx: {e}") from e

    from docx.shared import Pt
    label = f"{number.strip()}. {title.strip()}" if number.strip() else title.strip()
    try:
        doc.add_paragraph(label, style="Heading 1")
    except (KeyError, ValueError):
        h = doc.add_paragraph()
        r = h.add_run(label)
        r.bold = True
        r.font.size = Pt(14)
    if description.strip():
        doc.add_paragraph(description.strip())

    out = io.BytesIO()
    doc.save(out)
    base = Path(file.filename).stem or "document"
    return Response(
        content=out.getvalue(),
        media_type=DOCX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{base}.docx"',
            "X-Question-Added": "1",
            "Access-Control-Expose-Headers": "Content-Disposition, X-Question-Added",
        },
    )


# --- Document organizer (desktop-only) -------------------------------------

def _safe_dir(folder_str: str) -> Path:
    if not folder_str or not folder_str.strip():
        raise HTTPException(400, "Folder path required.")
    try:
        p = Path(folder_str).expanduser().resolve()
    except (OSError, RuntimeError) as e:
        raise HTTPException(400, f"Bad path: {e}") from e
    if not p.exists() or not p.is_dir():
        raise HTTPException(404, f"Folder not found: {p}")
    return p


def _list_docx(folder: Path) -> list[dict]:
    items: list[dict] = []
    try:
        for f in folder.iterdir():
            try:
                if not f.is_file():
                    continue
                if f.suffix.lower() != ".docx":
                    continue
                if f.name.startswith("~$"):  # Word lock files
                    continue
                st = f.stat()
                items.append({
                    "name": f.name,
                    "path": str(f),
                    "size": st.st_size,
                    "modified": st.st_mtime,
                })
            except OSError:
                continue
    except OSError:
        return []
    items.sort(key=lambda x: -x["modified"])
    return items


@app.get("/api/organizer/generated")
async def organizer_generated() -> dict:
    _require_desktop()
    try:
        GENERATED_DIR.mkdir(parents=True, exist_ok=True)
    except OSError as e:
        raise HTTPException(500, f"Could not create generated folder: {e}") from e
    return {"folder": str(GENERATED_DIR), "files": _list_docx(GENERATED_DIR)}


class ScanReq(BaseModel):
    folder: str


@app.post("/api/organizer/scan")
async def organizer_scan(body: ScanReq) -> dict:
    _require_desktop()
    p = _safe_dir(body.folder)
    return {"folder": str(p), "files": _list_docx(p)}


class OpenReq(BaseModel):
    path: str


@app.post("/api/organizer/open")
async def organizer_open(body: OpenReq) -> Response:
    _require_desktop()
    try:
        p = Path(body.path).expanduser().resolve()
    except (OSError, RuntimeError) as e:
        raise HTTPException(400, f"Bad path: {e}") from e
    if not p.exists() or not p.is_file():
        raise HTTPException(404, "File not found.")
    if p.suffix.lower() != ".docx":
        raise HTTPException(400, "Only .docx files.")
    try:
        data = p.read_bytes()
    except OSError as e:
        raise HTTPException(500, f"Could not read file: {e}") from e
    return Response(
        content=data,
        media_type=DOCX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{p.name}"',
            "X-File-Path": str(p),
            "Access-Control-Expose-Headers": "Content-Disposition, X-File-Path",
        },
    )


@app.post("/api/organizer/new-blank")
async def organizer_new_blank() -> Response:
    _require_desktop()
    doc = Document()
    doc.add_paragraph("")
    out = io.BytesIO()
    doc.save(out)
    return Response(
        content=out.getvalue(),
        media_type=DOCX_MIME,
        headers={
            "Content-Disposition": 'attachment; filename="untitled.docx"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


def _sanitize_filename(name: str) -> str:
    name = name.strip()
    name = "".join(c if c.isalnum() or c in "-_ ." else "_" for c in name)
    name = name.strip(". _")  # Windows hates trailing dot/space
    if not name:
        return ""
    if not name.lower().endswith(".docx"):
        name += ".docx"
    return name


def _unique_target(target_dir: Path, safe_name: str) -> Path:
    """Return a path in target_dir that won't clobber an existing file.

    'foo.docx' -> 'foo.docx', or 'foo-1.docx', 'foo-2.docx', ... if taken.
    Used for auto-saved template output so generating twice never silently
    overwrites an earlier record.
    """
    target = target_dir / safe_name
    if not target.exists():
        return target
    stem, suffix = target.stem, (target.suffix or ".docx")
    i = 1
    while True:
        cand = target_dir / f"{stem}-{i}{suffix}"
        if not cand.exists():
            return cand
        i += 1


@app.post("/api/organizer/save")
async def organizer_save(
    file: UploadFile = File(...),
    edits: str = Form("{}"),
    name: str = Form(...),
    folder: str = Form("generated"),  # "generated" | "current"
    current_path: str = Form(""),
) -> dict:
    _require_desktop()
    safe_name = _sanitize_filename(name)
    if not safe_name:
        raise HTTPException(400, "Filename can't be blank.")

    try:
        parsed_edits = json.loads(edits) if edits.strip() else {}
        if not isinstance(parsed_edits, dict):
            raise ValueError("edits must be a JSON object")
    except (ValueError, json.JSONDecodeError) as e:
        raise HTTPException(400, f"Invalid edits payload: {e}") from e

    raw = await file.read()
    if not raw:
        raise HTTPException(400, "Uploaded file is empty.")

    if parsed_edits:
        try:
            out_bytes, _applied = apply_block_edits_bytes(raw, parsed_edits)
        except Exception as e:
            raise HTTPException(500, f"Failed to apply edits: {e}") from e
    else:
        out_bytes = raw

    if folder == "current" and current_path.strip():
        try:
            cp = Path(current_path).expanduser().resolve()
        except (OSError, RuntimeError) as e:
            raise HTTPException(400, f"Bad current path: {e}") from e
        target_dir = cp.parent
        if not target_dir.exists() or not target_dir.is_dir():
            raise HTTPException(404, "Original folder no longer exists.")
    else:
        target_dir = GENERATED_DIR
        target_dir.mkdir(parents=True, exist_ok=True)

    target = target_dir / safe_name
    try:
        target.write_bytes(out_bytes)
    except OSError as e:
        raise HTTPException(500, f"Could not write file: {e}") from e

    return {
        "path": str(target),
        "name": target.name,
        "folder": str(target_dir),
        "size": len(out_bytes),
    }


# --- AI rewrite -------------------------------------------------------------

class RewriteIn(BaseModel):
    text: str
    instruction: str | None = ""
    target_words: int | None = None
    api_key: str | None = ""


def _ai_rewrite(text: str, instruction: str, target_words: int, api_key: str) -> str:
    """Single-shot Gemini call. Returns the rewritten text only."""
    prompt = (
        "Rewrite the following text. Hard constraints:\n"
        f"- Aim for ~{target_words} words (within +/- 15%). This preserves document layout.\n"
        "- Keep the technical meaning, register, and tone identical.\n"
        "- Plain prose only: no markdown, no quoting, no preamble like 'Here is'.\n"
        f"- Instruction from the user: {instruction or 'paraphrase faithfully'}\n\n"
        "Text:\n"
        f"{text}"
    )
    out, _ = gemini_call([{"text": prompt}], api_key=api_key, max_tokens=1024, timeout=45)
    return out or text


@app.post("/api/ai/rewrite")
async def ai_rewrite(body: RewriteIn, authorization: str = Header("")) -> dict:
    if not body.text.strip():
        raise HTTPException(400, "Text is empty.")
    key, charged = _paid_key(authorization, body.api_key or "", "rewrite")
    target = body.target_words or max(1, len(body.text.split()))
    try:
        new = _ai_rewrite(body.text, body.instruction or "", target, key)
    except urllib.error.HTTPError as e:
        _refund(charged, "rewrite")
        msg = e.read().decode("utf-8", errors="replace") if e.fp else str(e)
        raise HTTPException(502, f"Gemini API error {e.code}: {msg[:200]}") from e
    except Exception as e:
        _refund(charged, "rewrite")
        raise HTTPException(502, f"AI call failed: {e}") from e
    return {
        "text": new,
        "words_in": len(body.text.split()),
        "words_out": len(new.split()),
    }


# --- Template fill ----------------------------------------------------------

class TemplateFill(BaseModel):
    template: str = "neutral"
    course_title: str = ""
    register_number: str = ""
    ex_no: str = ""
    title: str = ""
    date: str = ""
    # formatting prefs (the pre-generation questionnaire)
    font: str = "Times New Roman"
    body_pt: int = 12
    margin: str = "normal"       # "normal" (1") | "narrow" (0.5")
    page_numbers: bool = True
    header_line: bool = True


@app.get("/api/templates")
async def list_templates() -> dict:
    """List built-in templates. Will expand once we have a templates_lib/."""
    return {
        "templates": [
            {
                "id": "neutral",
                "name": "Neutral record",
                "summary": "A4 · Times New Roman · header table · Aim → Result. Use when your teacher hasn't pinned a specific layout.",
                "fields": ["course_title", "register_number", "ex_no", "title", "date"],
            },
        ],
    }


@app.post("/api/template/fill")
async def fill_template(body: TemplateFill) -> Response:
    if body.template != "neutral":
        raise HTTPException(404, f"Unknown template: {body.template}")
    # Build the empty template to a temp file, then fill placeholders in-memory.
    from make_template import build_template

    with tempfile.TemporaryDirectory() as td:
        empty = Path(td) / "empty.docx"
        build_template(
            empty,
            font=(body.font.strip() or "Times New Roman")[:64],
            body_pt=min(max(body.body_pt, 9), 14),
            margin_in=0.5 if body.margin == "narrow" else 1.0,
            page_numbers=body.page_numbers,
            header_line=body.header_line,
        )
        doc = Document(str(empty))

    fills = {
        "{{COURSE_TITLE}}":    body.course_title,
        "{{REGISTER_NUMBER}}": body.register_number,
        "{{EX_NO}}":           body.ex_no,
        "{{TITLE}}":           body.title,
        "{{DATE}}":            body.date,
    }
    for find, repl in fills.items():
        replace_in_doc(doc, find, repl)

    out = io.BytesIO()
    doc.save(out)
    data = out.getvalue()
    safe = "".join(c if c.isalnum() or c in "-_ " else "_" for c in (body.title or "record")).strip()[:50] or "record"
    headers = {
        "Content-Disposition": f'attachment; filename="{safe}.docx"',
        "Access-Control-Expose-Headers": "Content-Disposition, X-File-Path",
    }
    # Desktop: persist template output to the Generated folder so it shows up in
    # the organizer, and hand the editor the on-disk path. Dedupe so generating
    # twice with the same title never silently overwrites an earlier record.
    if DESKTOP_MODE:
        try:
            GENERATED_DIR.mkdir(parents=True, exist_ok=True)
            saved = _unique_target(GENERATED_DIR, _sanitize_filename(f"{safe}.docx"))
            saved.write_bytes(data)
            headers["X-File-Path"] = str(saved)
        except OSError:
            pass  # non-fatal: fall back to open-in-editor only
    return Response(content=data, media_type=DOCX_MIME, headers=headers)


# --- AI code generation -------------------------------------------------------

class GenerateCodeIn(BaseModel):
    project_title: str
    question: str
    language: str = "python"
    mode: str = "auto"          # "auto" | "single" | "pair"
    extra_instructions: str = ""


_CODEGEN_SYSTEM = (
    "You write programs for a college Computer Networks lab record. "
    "Programs must be simple, deterministic, student-grade (clear variable "
    "names, a few comments, no over-engineering), and must run to completion "
    "without user interaction: hardcode sample inputs, never call input(). "
    "Socket programs must use 127.0.0.1, a port in 5000-5999, print a line "
    "when the server starts, and the client must terminate on its own so the "
    "run can complete."
)


def _ai_generate_code(body: GenerateCodeIn, api_key: str) -> dict:
    prompt = (
        f"Project/experiment context: {body.project_title}\n"
        f"Question to solve: {body.question}\n"
        f"Language: {body.language}\n"
        f"Mode: {body.mode} (auto = you decide; use pair ONLY for "
        f"client/server programs)\n"
        f"Extra instructions: {body.extra_instructions or 'none'}\n\n"
        "Respond with ONLY a JSON object, no markdown fences, in one of these "
        "shapes:\n"
        '{"mode":"single","code":"..."}\n'
        '{"mode":"pair","server_code":"...","client_code":"..."}\n'
    )
    text, _ = gemini_call(
        [{"text": prompt}],
        api_key=api_key,
        system=_CODEGEN_SYSTEM,
        max_tokens=4000,
        timeout=90,
        force_json=True,
    )
    return json.loads(strip_fences(text))


@app.post("/api/ai/generate-code")
async def ai_generate_code(body: GenerateCodeIn) -> dict:
    _require_desktop()
    key = _profile_key()
    if not key:
        raise HTTPException(400, "No Gemini API key configured. Add one in Profile.")
    if not body.question.strip():
        raise HTTPException(400, "Question is empty.")
    if body.language not in LANGS:
        raise HTTPException(400, f"language must be one of {LANGS}")
    try:
        result = _ai_generate_code(body, key)
    except urllib.error.HTTPError as e:
        msg = e.read().decode("utf-8", errors="replace") if e.fp else str(e)
        raise HTTPException(502, f"Gemini API error {e.code}: {msg[:200]}") from e
    except (json.JSONDecodeError, KeyError) as e:
        raise HTTPException(502, f"AI returned malformed code payload: {e}") from e
    except Exception as e:
        raise HTTPException(502, f"AI call failed: {e}") from e
    result["language"] = body.language
    return result


# --- Generate question (web-safe: Gemini writes code + output, no execution) --

_GEN_SYSTEM = (
    "You write programs and their terminal output for a college lab record. "
    "Programs must be simple, deterministic, student-grade (clear variable "
    "names, a few comments, no over-engineering), and self-contained: hardcode "
    "sample inputs, never read stdin. The output you write must be EXACTLY "
    "what running the program would print — no commentary, no markdown, under "
    "25 lines."
)


@app.post("/api/generate/question")
async def generate_question(
    file: UploadFile = File(...),
    number: str = Form(""),
    title: str = Form(...),
    question: str = Form(""),
    language: str = Form("python"),
    urk: str = Form(""),
    urk_mode: str = Form("line"),        # "line" | "title" | "prompt" | "none"
    include_code_shot: str = Form("1"),
    api_key: str = Form(""),
    authorization: str = Header(""),
) -> Response:
    """One generated question: AI code + simulated run, rendered screenshots,
    inserted into the uploaded record. The frontend loops this per question."""
    if not file.filename or not file.filename.lower().endswith(".docx"):
        raise HTTPException(400, "Please upload a .docx file.")
    if not title.strip():
        raise HTTPException(400, "Title is required.")
    if language not in LANGS:
        raise HTTPException(400, f"language must be one of {LANGS}")
    data = await file.read()
    if not data:
        raise HTTPException(400, "Uploaded file is empty.")

    key, charged = _paid_key(authorization, api_key, "gen_question")
    prompt = (
        f"Experiment/context: {title}\n"
        f"Task: {question.strip() or title}\n"
        f"Language: {language}\n\n"
        "Respond with ONLY a JSON object, no markdown fences:\n"
        '{"filename":"program' + _ext(language) + '","command":"...",'
        '"code":"...","output":"..."}\n'
        "code = the full program source. command = how a student runs it on "
        "Windows cmd. output = the exact terminal output of that run."
    )
    try:
        text, _ = gemini_call(
            [{"text": prompt}], api_key=key, system=_GEN_SYSTEM,
            max_tokens=4000, timeout=90, force_json=True,
        )
        payload = json.loads(strip_fences(text))
        code = payload["code"]
        output = payload.get("output", "")
        command = payload.get("command", "")
        filename = payload.get("filename", "program" + _ext(language))

        images: list[bytes] = []
        if include_code_shot == "1":
            images.append(render_code_shot(code, language=language, filename=filename))
        images.append(render_terminal_shot(
            output, command=command, cwd=r"C:\Users\{urk}\Desktop\Lab",
            style="cmd", urk=urk, urk_mode=urk_mode,
        ))
        out_bytes = add_solved_question_bytes(
            data, number=number, title=title, description=question,
            code=code, images=images,
        )
    except urllib.error.HTTPError as e:
        _refund(charged, "gen_question")
        msg = e.read().decode("utf-8", errors="replace") if e.fp else str(e)
        raise HTTPException(502, f"Gemini API error {e.code}: {msg[:200]}") from e
    except (json.JSONDecodeError, KeyError) as e:
        _refund(charged, "gen_question")
        raise HTTPException(502, f"AI returned a malformed payload: {e}") from e
    except Exception as e:
        _refund(charged, "gen_question")
        raise HTTPException(502, f"Generation failed: {e}") from e

    base = Path(file.filename).stem or "record"
    return Response(
        content=out_bytes,
        media_type=DOCX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{base}.docx"',
            "X-Question-Added": "1",
            "X-Images-Added": str(len(images)),
            "Access-Control-Expose-Headers":
                "Content-Disposition, X-Question-Added, X-Images-Added",
        },
    )


# --- Solve pipeline: run + screenshot + insert --------------------------------

@app.post("/api/document/add-solved-question")
async def doc_add_solved_question(
    file: UploadFile = File(...),
    number: str = Form(""),
    title: str = Form(...),
    description: str = Form(""),
    language: str = Form("python"),
    mode: str = Form("single"),          # "single" | "pair"
    code: str = Form(""),                # single mode
    server_code: str = Form(""),         # pair mode
    client_code: str = Form(""),         # pair mode
    stdin_text: str = Form(""),
    filename: str = Form(""),            # display name for prompt/tab, e.g. q1.py
    urk: str = Form(""),
    urk_mode: str = Form("line"),        # "line" | "title" | "prompt" | "none"
    terminal_style: str = Form("cmd"),   # "cmd" | "powershell"
    include_code_shot: str = Form("0"),  # "1" -> also render VS Code editor image
    cwd: str = Form(r"C:\Users\{urk}\Desktop\Lab"),
    timeout: float = Form(15.0),
) -> Response:
    if not file.filename or not file.filename.lower().endswith(".docx"):
        raise HTTPException(400, "Please upload a .docx file.")
    if not title.strip():
        raise HTTPException(400, "Title is required.")
    if language not in LANGS:
        raise HTTPException(400, f"language must be one of {LANGS}")
    data = await file.read()
    if not data:
        raise HTTPException(400, "Uploaded file is empty.")

    images: list[bytes] = []
    ran_ok = True
    try:
        if mode == "pair":
            if not server_code.strip() or not client_code.strip():
                raise HTTPException(400, "Pair mode needs server_code and client_code.")
            srv, cli = run_pair(server_code, client_code, language=language,
                                client_stdin=stdin_text, client_timeout=timeout)
            ran_ok = cli.ok
            if include_code_shot == "1":
                images.append(render_code_shot(server_code, language=language,
                                               filename="server" + _ext(language)))
            images.append(render_terminal_shot(
                srv.output, command=srv.command, cwd=cwd, style=terminal_style,
                urk=urk, urk_mode=("title" if urk_mode != "none" else "none"),
                trailing_prompt=False))
            if include_code_shot == "1":
                images.append(render_code_shot(client_code, language=language,
                                               filename="client" + _ext(language)))
            images.append(render_terminal_shot(
                cli.output, command=cli.command, cwd=cwd, style=terminal_style,
                urk=urk, urk_mode=urk_mode))
            record_code = (server_code + "\n\n" + client_code) if not code.strip() else code
        else:
            if not code.strip():
                raise HTTPException(400, "Single mode needs code.")
            res = run_single(code, language=language, stdin_text=stdin_text,
                             timeout=timeout, filename=filename or None)
            ran_ok = res.ok
            if include_code_shot == "1":
                images.append(render_code_shot(code, language=language,
                                               filename=filename or ("program" + _ext(language))))
            images.append(render_terminal_shot(
                res.output, command=res.command, cwd=cwd, style=terminal_style,
                urk=urk, urk_mode=urk_mode))
            record_code = code
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Run/render failed: {e}") from e

    try:
        out_bytes = add_solved_question_bytes(
            data, number=number, title=title, description=description,
            code=record_code, images=images,
        )
    except Exception as e:
        raise HTTPException(500, f"Failed to write question into docx: {e}") from e

    base = Path(file.filename).stem or "document"
    return Response(
        content=out_bytes,
        media_type=DOCX_MIME,
        headers={
            "Content-Disposition": f'attachment; filename="{base}.docx"',
            "X-Question-Added": "1",
            "X-Run-Ok": "1" if ran_ok else "0",
            "X-Images-Added": str(len(images)),
            "Access-Control-Expose-Headers":
                "Content-Disposition, X-Question-Added, X-Run-Ok, X-Images-Added",
        },
    )


def _ext(language: str) -> str:
    return {"python": ".py", "c": ".c", "cpp": ".cpp", "java": ".java"}[language]


# Static files (CSS/JS) — must be after explicit routes
app.mount("/static", StaticFiles(directory=STATIC), name="static")
