"""Verify the generated template opens cleanly and the find/replace pipeline
can fill its placeholders end-to-end.
"""

from pathlib import Path
from docx import Document

from autolab import iter_all_paragraphs, replace_in_docx

TEMPLATE = Path(__file__).parent / "template_empty.docx"
FILLED = Path(__file__).parent / "template_filled_test.docx"

FIELDS = {
    "{{COURSE_TITLE}}":    "Computer Networks Lab",
    "{{REGISTER_NUMBER}}": "URK24CS1021",
    "{{EX_NO}}":           "1",
    "{{TITLE}}":           "Basic Network Troubleshooting Commands",
    "{{DATE}}":            "13.06.2026",
}


def count_all(path):
    doc = Document(str(path))
    text = "\n".join(p.text for p in iter_all_paragraphs(doc))
    return {k: text.count(k) for k in FIELDS}


# Step 1: read template back, count placeholders
before = count_all(TEMPLATE)
print("Placeholders in empty template:")
for k, v in before.items():
    print(f"  {k:24} -> {v}")

# Step 2: copy template -> filled via N replace passes
import shutil
shutil.copyfile(TEMPLATE, FILLED)
totals = {}
for find, repl in FIELDS.items():
    totals[find] = replace_in_docx(FILLED, FILLED, find, repl)

print("\nReplacements performed:")
for k, v in totals.items():
    print(f"  {k:24} -> {v}")

# Step 3: verify nothing left
remaining = count_all(FILLED)
print("\nRemaining placeholders after fill (should all be 0):")
for k, v in remaining.items():
    print(f"  {k:24} -> {v}")

ok = all(v == 0 for v in remaining.values()) and all(b >= 1 for b in before.values())
print("\nPASS" if ok else "FAIL")
