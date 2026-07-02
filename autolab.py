"""AutoLAB core — docx find/replace engine.

Pure functions. No UI. The web server (server.py) and the smoke tests both
import from here.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.table import _Cell
from docx.text.paragraph import Paragraph


def replace_in_paragraph(paragraph: Paragraph, find: str, replace: str) -> int:
    """Replace every occurrence of `find` with `replace` inside one paragraph.

    Word stores paragraph text as a sequence of runs (<w:r>) so a literal like
    "URK24CS1021" can be split across runs and a naive per-run replace misses
    it. We try the fast per-run path first, then fall back to consolidating
    all run text into the first run for any remaining occurrences. Formatting
    on the first run is preserved; secondary runs lose their text.
    """
    n = paragraph.text.count(find)
    if n == 0:
        return 0

    for run in paragraph.runs:
        if find in run.text:
            run.text = run.text.replace(find, replace)

    if find not in paragraph.text:
        return n

    full = "".join(r.text for r in paragraph.runs)
    if find in full and paragraph.runs:
        paragraph.runs[0].text = full.replace(find, replace)
        for r in paragraph.runs[1:]:
            r.text = ""
    return n


def _iter_cell_paragraphs(cell: _Cell) -> Iterable[Paragraph]:
    for p in cell.paragraphs:
        yield p
    for table in cell.tables:
        for row in table.rows:
            for c in row.cells:
                yield from _iter_cell_paragraphs(c)


def iter_all_paragraphs(doc: DocumentType) -> Iterable[Paragraph]:
    """Yield every paragraph: body, tables, nested tables, headers, footers."""
    for p in doc.paragraphs:
        yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_cell_paragraphs(cell)
    for section in doc.sections:
        for hf in (section.header, section.footer):
            for p in hf.paragraphs:
                yield p
            for table in hf.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from _iter_cell_paragraphs(cell)


def replace_in_doc(doc: DocumentType, find: str, replace: str) -> int:
    return sum(replace_in_paragraph(p, find, replace) for p in iter_all_paragraphs(doc))


def replace_in_docx(src: str | Path, dst: str | Path, find: str, replace: str) -> int:
    doc = Document(str(src))
    n = replace_in_doc(doc, find, replace)
    doc.save(str(dst))
    return n


def replace_in_docx_bytes(data: bytes, find: str, replace: str) -> tuple[bytes, int]:
    """Bytes-in, bytes-out variant — what the web server calls."""
    doc = Document(io.BytesIO(data))
    n = replace_in_doc(doc, find, replace)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), n


# --- Block-level editing -----------------------------------------------------

def _block_path(*parts: object) -> str:
    return ".".join(str(p) for p in parts)


def extract_blocks(doc: DocumentType) -> list[dict]:
    """Walk a Document and produce stable addressable blocks for the editor.

    Each block has:
      id     — path-based, stable across reads of the same doc
      kind   — paragraph | cell | header | footer
      text   — current text content
      style  — paragraph style name (e.g. 'Heading 1')
      loc    — human-readable location for the UI
      words  — current word count (used as the rewrite budget)
      chars  — current char count
    """
    blocks: list[dict] = []

    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "":
            continue
        blocks.append(_block_for(p, _block_path("body", "p", i), kind="paragraph", loc=f"Body · ¶{i + 1}"))

    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                for pi, p in enumerate(cell.paragraphs):
                    if p.text.strip() == "":
                        continue
                    blocks.append(_block_for(
                        p,
                        _block_path("body", "tbl", ti, "r", ri, "c", ci, "p", pi),
                        kind="cell",
                        loc=f"Table {ti + 1} · R{ri + 1}C{ci + 1}",
                    ))

    for si, section in enumerate(doc.sections):
        for pi, p in enumerate(section.header.paragraphs):
            if p.text.strip() == "":
                continue
            blocks.append(_block_for(
                p, _block_path("sec", si, "hdr", "p", pi), kind="header",
                loc=f"Header · §{si + 1} ¶{pi + 1}"))
        for pi, p in enumerate(section.footer.paragraphs):
            if p.text.strip() == "":
                continue
            blocks.append(_block_for(
                p, _block_path("sec", si, "ftr", "p", pi), kind="footer",
                loc=f"Footer · §{si + 1} ¶{pi + 1}"))
        # tables in header/footer
        for hi, hf in enumerate((section.header, section.footer)):
            tag = "hdr" if hi == 0 else "ftr"
            for ti, table in enumerate(hf.tables):
                for ri, row in enumerate(table.rows):
                    for ci, cell in enumerate(row.cells):
                        for pi, p in enumerate(cell.paragraphs):
                            if p.text.strip() == "":
                                continue
                            blocks.append(_block_for(
                                p,
                                _block_path("sec", si, tag, "tbl", ti, "r", ri, "c", ci, "p", pi),
                                kind="cell",
                                loc=f"{'Header' if hi == 0 else 'Footer'} Table § {si + 1}",
                            ))

    return blocks


def _block_for(p: Paragraph, path: str, *, kind: str, loc: str) -> dict:
    text = p.text
    return {
        "id": path,
        "kind": kind,
        "text": text,
        "style": p.style.name if p.style is not None else "Normal",
        "loc": loc,
        "words": len(text.split()),
        "chars": len(text),
    }


def _resolve_paragraph(doc: DocumentType, path: str) -> Paragraph | None:
    parts = path.split(".")
    it = iter(parts)
    try:
        scope = next(it)
        if scope == "body":
            tag = next(it)
            if tag == "p":
                idx = int(next(it))
                return doc.paragraphs[idx]
            if tag == "tbl":
                t_idx = int(next(it)); assert next(it) == "r"
                r_idx = int(next(it)); assert next(it) == "c"
                c_idx = int(next(it)); assert next(it) == "p"
                p_idx = int(next(it))
                return doc.tables[t_idx].rows[r_idx].cells[c_idx].paragraphs[p_idx]
        elif scope == "sec":
            s_idx = int(next(it))
            section = doc.sections[s_idx]
            place = next(it)
            container = section.header if place == "hdr" else section.footer
            tag = next(it)
            if tag == "p":
                p_idx = int(next(it))
                return container.paragraphs[p_idx]
            if tag == "tbl":
                t_idx = int(next(it)); assert next(it) == "r"
                r_idx = int(next(it)); assert next(it) == "c"
                c_idx = int(next(it)); assert next(it) == "p"
                p_idx = int(next(it))
                return container.tables[t_idx].rows[r_idx].cells[c_idx].paragraphs[p_idx]
    except (StopIteration, AssertionError, IndexError, ValueError):
        return None
    return None


def _set_paragraph_text(paragraph: Paragraph, new_text: str) -> None:
    """Replace the entire text of a paragraph, preserving the first run's format."""
    if paragraph.text == new_text:
        return
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def apply_block_edits_bytes(data: bytes, edits: dict[str, str]) -> tuple[bytes, int]:
    """Apply id -> new text edits to a docx, return new bytes + applied count."""
    doc = Document(io.BytesIO(data))
    applied = 0
    for block_id, new_text in edits.items():
        p = _resolve_paragraph(doc, block_id)
        if p is None:
            continue
        if p.text == new_text:
            continue
        _set_paragraph_text(p, new_text)
        applied += 1
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue(), applied
