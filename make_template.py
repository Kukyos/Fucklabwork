"""AutoLAB template generator: produce an empty 'global' lab-record .docx.

The output is a neutral, teacher-friendly skeleton: A4 + 1" margins, Times
New Roman, bordered header table, page numbers, register number in both
header and footer. Placeholders use {{NAME}} so the find/replace pipeline in
autolab.py can fill them in later.

Usage:
    python make_template.py [output.docx]
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


FONT = "Times New Roman"
BODY_PT = 12
HEADING_PT = 14

SECTIONS = [
    ("Aim", "[Write the aim of the experiment here.]"),
    ("Algorithm / Procedure", "[Outline the steps or algorithm here.]"),
    ("Program", "[Paste the program / source code here.]"),
    ("Output", "[Paste the output or screenshot here.]"),
    ("Result", "[State the final result here.]"),
]


def _set_rfonts(rpr, font_name=FONT):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), font_name)


def _set_cell_border(cell, color="000000", size=6):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), color)
        borders.append(b)


def _insert_field(paragraph, instr, cached="1"):
    """Append a Word field (e.g. PAGE, NUMPAGES) to a paragraph in doc order."""
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), instr)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    _set_rfonts(rpr)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = cached
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)


def _style_styles(doc, font=FONT, body_pt=BODY_PT):
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(body_pt)
    _set_rfonts(normal.element.get_or_add_rPr(), font)

    for name, size in (("Heading 1", body_pt + 2), ("Heading 2", body_pt + 1)):
        s = doc.styles[name]
        s.font.name = font
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        _set_rfonts(s.element.get_or_add_rPr(), font)


def _setup_page(section, margin_in=1.0):
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    for attr in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attr, Inches(margin_in))


def _build_header(section):
    p = section.header.paragraphs[0]
    # Right tab stop near the right margin (A4 content width ~ 6.27" ~ 15.9cm)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("{{COURSE_TITLE}}").bold = True
    p.add_run("\t")
    p.add_run("{{REGISTER_NUMBER}}").bold = True


def _build_footer(section, page_numbers=True):
    p = section.footer.paragraphs[0]
    p.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
    p.add_run("{{REGISTER_NUMBER}}")
    if page_numbers:
        p.add_run("\tPage ")
        _insert_field(p, "PAGE")
        p.add_run(" of ")
        _insert_field(p, "NUMPAGES")


def _build_top_table(doc):
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Column widths: narrow ID column + wide title column. Total ~ 6.27" content width.
    widths = (Inches(2.25), Inches(4.0))

    rows = [
        ("Ex.No: {{EX_NO}}", "Title: {{TITLE}}"),
        ("Date: {{DATE}}",   "{{REGISTER_NUMBER}}"),
    ]
    for r, row_data in enumerate(rows):
        for c, text in enumerate(row_data):
            cell = table.rows[r].cells[c]
            cell.width = widths[c]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.bold = True
            _set_cell_border(cell)


def _add_section_heading(doc, text):
    p = doc.add_paragraph(text, style="Heading 1")
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p


def _add_placeholder(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_after = Pt(6)
    return p


def build_template(out_path: Path, *, font: str = FONT, body_pt: int = BODY_PT,
                   margin_in: float = 1.0, page_numbers: bool = True,
                   header_line: bool = True) -> None:
    doc = Document()
    _style_styles(doc, font, body_pt)
    section = doc.sections[0]
    _setup_page(section, margin_in)
    if header_line:
        _build_header(section)
    _build_footer(section, page_numbers)
    _build_top_table(doc)
    doc.add_paragraph()  # spacer
    for name, placeholder in SECTIONS:
        _add_section_heading(doc, name)
        _add_placeholder(doc, placeholder)
    doc.save(str(out_path))


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent / "template_empty.docx"
    build_template(out)
    print(f"Wrote: {out}")
