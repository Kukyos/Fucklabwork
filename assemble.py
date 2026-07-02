"""AutoLAB assemble — write a solved question into a .docx.

Per question, in order:
    N. Title                    (Heading 1 when the doc has that style)
    description                 (optional)
    Program:                    (bold label)
    <code as a shaded monospace block — real text, copy-pasteable>
    Output:                     (bold label)
    <screenshot PNGs from labshot, centered, sized to content width>

Follows the same conventions as autolab.py: pure functions, a doc-level API
plus a bytes-in/bytes-out wrapper for the web server.
"""

from __future__ import annotations

import io
from typing import Iterable

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CODE_FONT = "Consolas"
CODE_PT = 10
# A4 with 1" margins leaves ~6.27" of content width; 6.0" gives breathing room.
DEFAULT_IMAGE_WIDTH_IN = 6.0


def _shade_paragraph(paragraph, fill: str = "F2F2F2") -> None:
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def _set_code_font(run) -> None:
    run.font.name = CODE_FONT
    run.font.size = Pt(CODE_PT)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), CODE_FONT)


def add_heading(doc: DocumentType, number: str, title: str) -> None:
    label = f"{number.strip()}. {title.strip()}" if number.strip() else title.strip()
    try:
        doc.add_paragraph(label, style="Heading 1")
    except (KeyError, ValueError):
        p = doc.add_paragraph()
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(14)


def add_label(doc: DocumentType, text: str) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)


def add_code_block(doc: DocumentType, code: str, *, shaded: bool = True) -> None:
    """One paragraph, one run per line separated by soft breaks. Keeps the
    block together visually and survives block-editor round trips (it is a
    single paragraph, so extract_blocks sees one block, not 40)."""
    lines = code.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    while lines and lines[-1].strip() == "":
        lines.pop()
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.space_after = Pt(6)
    if shaded:
        _shade_paragraph(p)
    for i, line in enumerate(lines):
        run = p.add_run(line if line else " ")
        _set_code_font(run)
        run.font.color.rgb = RGBColor(0x20, 0x20, 0x20)
        if i < len(lines) - 1:
            run.add_break()


def add_images(doc: DocumentType, images: Iterable[bytes],
               *, width_in: float = DEFAULT_IMAGE_WIDTH_IN) -> None:
    for png in images:
        doc.add_picture(io.BytesIO(png), width=Inches(width_in))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[-1].paragraph_format.space_after = Pt(8)


def add_solved_question(
    doc: DocumentType,
    *,
    number: str = "",
    title: str,
    description: str = "",
    code: str = "",
    images: Iterable[bytes] = (),
    code_label: str = "Program:",
    output_label: str = "Output:",
    code_shaded: bool = True,
    image_width_in: float = DEFAULT_IMAGE_WIDTH_IN,
) -> None:
    add_heading(doc, number, title)
    if description.strip():
        doc.add_paragraph(description.strip())
    if code.strip():
        if code_label:
            add_label(doc, code_label)
        add_code_block(doc, code, shaded=code_shaded)
    images = list(images)
    if images:
        if output_label:
            add_label(doc, output_label)
        add_images(doc, images, width_in=image_width_in)


def add_solved_question_bytes(data: bytes, **kwargs) -> bytes:
    """Bytes-in, bytes-out variant — what the web server calls."""
    doc = Document(io.BytesIO(data))
    add_solved_question(doc, **kwargs)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
