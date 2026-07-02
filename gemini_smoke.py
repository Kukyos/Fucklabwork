"""Offline check for gemini.py — mocks the API, verifies docx image swap."""

import io
import zipfile

from PIL import Image

import gemini


def _docx_with_image() -> bytes:
    from docx import Document
    from docx.shared import Inches

    import os
    img = io.BytesIO()
    # noisy image so the PNG stays above the icon-skip size threshold
    Image.frombytes("RGB", (300, 100), os.urandom(300 * 100 * 3)).save(img, "PNG")
    img.seek(0)
    doc = Document()
    doc.add_paragraph("URK24CS1021 record")
    doc.add_picture(img, width=Inches(2))
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def main() -> None:
    assert gemini.strip_fences('```json\n{"a":1}\n```') == '{"a":1}'
    assert gemini.strip_fences('{"a":1}') == '{"a":1}'

    replacement = io.BytesIO()
    Image.new("RGB", (300, 100), "white").save(replacement, "PNG")
    calls = []

    def fake_call(parts, **kw):
        calls.append(kw.get("model", gemini.TEXT_MODEL))
        if kw.get("want_image"):
            return "", replacement.getvalue()
        return "YES", None

    real = gemini.gemini_call
    gemini.gemini_call = fake_call
    try:
        src = _docx_with_image()
        out, patched, scanned = gemini.patch_docx_images(src, "URK24CS1021", "URK24CS9999", "fake-key")
    finally:
        gemini.gemini_call = real

    assert scanned == 1 and patched == 1, (patched, scanned)
    assert len(calls) == 2  # vision check + image edit
    zin = zipfile.ZipFile(io.BytesIO(out))
    media = [n for n in zin.namelist() if n.startswith("word/media/")]
    assert media and zin.read(media[0]) == replacement.getvalue()
    # docx still opens
    from docx import Document
    Document(io.BytesIO(out))
    print("PASS")


if __name__ == "__main__":
    main()
