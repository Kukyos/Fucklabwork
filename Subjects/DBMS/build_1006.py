"""One-off: assemble URK24CS1006's Ex1B + Ex2 docx from screenshots he
already captured himself (MySQL, not our usual psql pipeline) -- we don't
run any SQL here, just place his real images + the commands visible in them
into the same record layout dbms_lab.py uses for CS1021.

Ex1B Q1 has no screenshot (he asked to leave it blank -- question text only).
"""
import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from dbms_lab import (COURSE, FONT, IMG_W, EX1B, EX2, _rfonts, _cell_border,
                       _label, _body)
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm
from assemble import add_code_block

HERE = Path(__file__).parent
URK = "URK24CS1006"
DATE = "29/07/2026"

QTEXT_1B = {n: qt for n, qt, _ in EX1B["questions"]}
QTEXT_2 = {n: qt for n, qt, _ in EX2["questions"]}

SHOTS_1B = HERE / "EX1bUKR24CS1006" / "ex1 output dbms"
RESULTS_1B = [(1, QTEXT_1B[1], None, None)] + [
    (n, QTEXT_1B[n], code, (SHOTS_1B / f"q{n}.png").read_bytes()) for n, code in [
        (2, "UPDATE UsersURK24CS1006 SET Email = 'newemail@example.com' WHERE ID = 123;"),
        (3, "DELETE FROM UsersURK24CS1006 WHERE Email = 'example@example.com';"),
        (4, "INSERT INTO EventURK24CS1006 (EventID, Name, EventDate, EventTime, VenueID, Description)\n"
            "VALUES (456, 'Music Concert', '2026-08-15', '2026-08-15 18:30:00', 1, 'Live music concert');"),
        (5, "UPDATE EventURK24CS1006 SET Description = 'Updated event description'\nWHERE EventID = 456;"),
        (6, "GRANT SELECT ON dbmslab.UsersURK24CS1006 TO 'john'@'localhost';"),
        (7, "REVOKE INSERT ON dbmslab.EventURK24CS1006 FROM 'mary'@'localhost';"),
        (8, "GRANT ALL PRIVILEGES ON dbmslab.TicketURK24CS1006 TO 'jane'@'localhost';"),
        (9, "GRANT SELECT ON dbmslab.* TO 'guest'@'localhost';\n"
            "REVOKE ALL PRIVILEGES ON dbmslab.* FROM 'guest'@'localhost';"),
        (10, "GRANT SELECT ON dbmslab.* TO 'guest'@'localhost';\n"
             "REVOKE ALL PRIVILEGES ON dbmslab.* FROM 'guest'@'localhost';"),
        (11, "COMMIT;"),
        (12, "ROLLBACK TO SAVEPOINT sp1;"),
        (13, "SAVEPOINT sp1;"),
        (14, "SET autocommit = 1;"),
        (15, "SET autocommit = 0;"),
    ]
]

SHOTS_2 = HERE / "ex2urk24cs1006" / "ex2 output dbms"
_SHOT_FILES_2 = sorted((SHOTS_2).glob("*.png"))  # filenames sort by timestamp = question order
assert len(_SHOT_FILES_2) == 15, len(_SHOT_FILES_2)
CODE_2 = [
    "use dbmslab\n"
    "INSERT INTO LocationURK24CS1006 (VenueID, Name, Address, City, State, Country)\n"
    "VALUES (1, 'Madison Square Garden', '4 Pennsylvania Plaza', 'New York', 'New York', 'USA');\n"
    "INSERT INTO UsersURK24CS1006 (ID, Name, Email, Password, Phone)\n"
    "VALUES (101, 'John Doe', 'john@example.com', 'password123', '9876543210');\n"
    "INSERT INTO EventURK24CS1006 (EventID, Name, EventDate, EventTime, VenueID, Description)\n"
    "VALUES (201, 'Rock Concert', '2023-08-15', '2023-08-15 19:00:00', 1, 'Live Rock Music');\n"
    "INSERT INTO TicketURK24CS1006 (TicketID, EventID, UserID, Price, Status, Barcode)\n"
    "VALUES (301, 201, 101, 100.00, 'Confirmed', 'BC301');",
    "SELECT * FROM TicketURK24CS1006;\nSELECT * FROM UsersURK24CS1006;\n"
    "SELECT * FROM EventURK24CS1006;\nSELECT * FROM LocationURK24CS1006;",
    "SELECT *\nFROM LocationURK24CS1006\nORDER BY City ASC, Country ASC;",
    "SELECT COUNT(*) AS TotalUsers\nFROM UsersURK24CS1006;",
    "SELECT *\nFROM EventURK24CS1006\nORDER BY EventDate DESC, EventTime DESC;",
    "SELECT EventID, AVG(Price) AS AveragePrice\nFROM TicketURK24CS1006\nGROUP BY EventID;",
    "SELECT EventID, SUM(Price) AS TotalPrice, COUNT(*) AS TicketCount\nFROM TicketURK24CS1006\nGROUP BY EventID;",
    "SELECT EventID, COUNT(*) AS NumberOfTickets\nFROM TicketURK24CS1006\nGROUP BY EventID;",
    "SELECT DISTINCT *\nFROM UsersURK24CS1006;",
    "SELECT *\nFROM EventURK24CS1006\nWHERE EventDate > '2023-07-30';",
    "SELECT SUM(Price) AS TotalConfirmedPrice\nFROM TicketURK24CS1006\nWHERE Status = 'Confirmed';",
    "SELECT CONCAT(Name, ' - ', Email) AS UserDetails\nFROM UsersURK24CS1006;",
    "SELECT *\nFROM LocationURK24CS1006\nWHERE City = 'New York' AND Country = 'USA';",
    "SELECT City, Country\nFROM LocationURK24CS1006\nORDER BY Country ASC;",
    "DELETE FROM TicketURK24CS1006\nWHERE Status = 'Cancelled';",
]
RESULTS_2 = [(n, QTEXT_2[n], CODE_2[n - 1], _SHOT_FILES_2[n - 1].read_bytes())
             for n in range(1, 16)]


def build_docx(spec: dict, results: list[tuple], out_path: Path) -> None:
    title = spec["title"].upper()
    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = FONT; normal.font.size = Pt(12)
    _rfonts(normal.element)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1.0))

    h = sec.header.paragraphs[0]
    from docx.enum.text import WD_TAB_ALIGNMENT
    h.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
    r = h.add_run(COURSE); r.bold = True
    h.add_run("\t"); r2 = h.add_run(URK); r2.bold = True
    pbdr = h._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "12")
    bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), "000000")
    bd.append(bot); pbdr.append(bd)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = f.add_run(f"{title}: {spec['name']}"); fr.bold = True

    t = doc.add_table(rows=2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    rows = [(f"Ex. No. {spec['title'][2:]}", spec["name"]), ("Date of Exercise", DATE)]
    widths = (Inches(1.9), Inches(4.35))
    for ri, (a, b) in enumerate(rows):
        for ci, txt in enumerate((a, b)):
            c = t.rows[ri].cells[ci]; c.width = widths[ci]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = p.add_run(txt); cr.font.name = FONT; cr.font.size = Pt(12)
            cr.bold = (ci == 0 or ri == 0)
            _cell_border(c)
    doc.add_paragraph()

    _label(doc, "Aim"); _body(doc, spec["aim"])
    _label(doc, "Description"); _body(doc, spec["desc"])

    for n, qtext, code, png in results:
        _label(doc, f"{n}. {qtext}")
        if code and png:
            _label(doc, "Sample Code:", size=12, space_before=2)
            add_code_block(doc, code)
            _label(doc, "Sample Output:", size=12, space_before=2)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            p.add_run().add_picture(io.BytesIO(png), width=Inches(IMG_W))

    _label(doc, "RESULT")
    _body(doc, "The SQL queries were executed and the desired output was obtained.")
    doc.save(str(out_path))


def main():
    for spec, results, name in [(EX1B, RESULTS_1B, "Ex1B"), (EX2, RESULTS_2, "Ex2")]:
        outdir = HERE / "output1006" / name
        shots = outdir / "screenshots"; shots.mkdir(parents=True, exist_ok=True)
        for n, _, _, png in results:
            if png:
                (shots / f"q{n:02d}.png").write_bytes(png)
        docx_path = outdir / f"{name}_{URK}.docx"
        build_docx(spec, results, docx_path)
        print("wrote", docx_path)


if __name__ == "__main__":
    main()
