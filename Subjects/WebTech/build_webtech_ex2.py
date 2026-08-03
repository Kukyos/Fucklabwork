"""Build Cleo's WebTech Ex2 record -- same pattern as build_webtech_ex1.py:
fills mam's Lab Record Format.docx template with the real HTML5 form page,
the real code, and real browser screenshots (via claude-in-chrome).

    python Subjects/WebTech/build_webtech_ex2.py

index.html + the output_*.jpg screenshots must already exist in output/Ex2/
(written by hand + captured via the browser, not by this script) -- this
script only assembles the docx around them.
"""
from __future__ import annotations

import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from assemble import add_code_block

HERE = Path(__file__).parent
MATERIALS = HERE / "materials"
OUT = HERE / "output" / "Ex2"
URK = "URK24CS1021"
COURSE = "23CS2049 - Web Technology Lab"
TITLE = "Registration Form - Join Respawn Point"
DATE = "27/07/2026"
GITHUB_LINK = "https://kukyos.github.io/WebTechLab/Ex2/index.html"
IMG_W = 6.2

AIM = ("To design and implement an HTML5 registration form using the various "
       "form input elements and controls, applied to the online gaming "
       "community theme.")

DESC = ("HTML5 introduced several new <input> types beyond plain text, each "
        "giving the browser its own built-in validation/UI: email (checks for "
        "an @ address), password (masks input), number (spinner + min/max), "
        "date (a date picker), range (a slider between min/max), and color (a "
        "color picker). Radio buttons let the user pick exactly one option "
        "from a group sharing the same name; checkboxes allow multiple. A "
        "<select> gives a fixed dropdown of options, while a <datalist> "
        "attached to a text input via the list attribute gives suggestions "
        "the user can still type past. Each <label>'s for attribute ties it to "
        "its input's id, and the fields are separated with <br> tags, the same "
        "way as in the sample form. The background and text colors are set "
        "with inline CSS, carried over from the Ex1 page.")

PROCEDURE = (
    "1. Reused the sample registration form as a reference for which input "
    "types were required (text, email, password, number, date, radio, "
    "checkbox, select, datalist, range, color).\n"
    "2. Reskinned every field to the Respawn Point gaming theme instead of "
    "the generic student-registration wording (Gamer Tag instead of Full "
    "Name, Main Platform instead of Gender, Games You're Into instead of "
    "Skills, and so on).\n"
    "3. Laid the fields out the same way as the sample form: a <label> with a "
    "for attribute, a <br>, the input, then <br><br> before the next field, "
    "with an HTML comment naming the input type above each one.\n"
    "4. Used the same background and text colors as the Ex1 page, applied "
    "with inline CSS on the <body> tag and on the labels, so this page looks "
    "like part of the same site.\n"
    "5. Added a link at the top back to the Ex1 homepage.\n"
    "6. Opened the page in the browser and filled in every field (typed the "
    "text fields, entered a date, picked a radio button and two checkboxes, "
    "chose a dropdown option, typed a datalist suggestion and moved the range "
    "slider) to confirm they all actually work, then took the output "
    "screenshots of the empty and the filled-in form."
)

RESULT = ("The HTML5 registration form was successfully designed and "
          "executed using the required input elements and controls. The "
          "output matched the expected layout and content.")


def set_placeholder(doc: Document, marker: str, text: str) -> None:
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.text.strip() == marker:
            ph = paras[i + 1]
            for r in list(ph.runs):
                r.text = ""
            if ph.runs:
                ph.runs[0].text = text
            else:
                ph.add_run(text)
            return
    raise ValueError(f"placeholder for {marker!r} not found")


def clear_placeholder(doc: Document, marker: str) -> "docx.text.paragraph.Paragraph":
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.text.strip() == marker:
            ph = paras[i + 1]
            for r in list(ph.runs):
                r.text = ""
            return ph
    raise ValueError(f"placeholder for {marker!r} not found")


def fill_header_regno(doc: Document, urk: str) -> None:
    """See build_webtech_ex1.py's copy of this function for why: the header's
    course-title cell is a Word content control the table API can't see, but
    the [regno] cell next to it is plain, so walk the header XML directly."""
    from docx.oxml.ns import qn
    header = doc.sections[0].header._element
    for t in header.iter(qn("w:t")):
        if t.text == "[regno]":
            t.text = urk
            return
    raise ValueError("[regno] placeholder not found in header")


def fill_footer_title(doc: Document, ex_label: str, title: str) -> None:
    """See build_webtech_ex1.py's copy: updates both the cached 'Ex. N' run
    text and the plain ' | Title of exercise' run, walking w:t directly."""
    from docx.oxml.ns import qn
    footer = doc.sections[0].footer._element
    found_ex, found_title = False, False
    for t in footer.iter(qn("w:t")):
        if t.text == "Ex. 1":
            t.text = ex_label
            found_ex = True
        elif t.text and t.text.strip() == "| Title of exercise":
            t.text = f" | {title}"
            found_title = True
    if not (found_ex and found_title):
        raise ValueError(f"footer placeholders not found (ex={found_ex}, title={found_title})")


def set_ex_number(docx_path: Path, label: str) -> None:
    """See build_webtech_ex1.py's copy: 'Ex. N' is bound to the extended
    'Company' property, patched post-save since python-docx has no API for it."""
    import re, zipfile
    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                text = data.decode("utf8")
                text = re.sub(r"<Company>.*?</Company>", f"<Company>{label}</Company>", text)
                data = text.encode("utf8")
            zout.writestr(item, data)
    tmp.replace(docx_path)


def build() -> Path:
    tmpl = next(MATERIALS.glob("Lab Record Format.docx"))
    doc = Document(str(tmpl))

    # see build_webtech_ex1.py -- the header's bound Title property was
    # stale ("MERN Full Stack...") from whoever the template was copied
    # from originally; fix it at the property level, not just the cached text
    doc.core_properties.title = COURSE
    doc.core_properties.author = URK
    doc.core_properties.last_modified_by = URK
    fill_header_regno(doc, URK)
    fill_footer_title(doc, "Ex. 2", TITLE)

    t = doc.tables[0]
    t.rows[0].cells[0].paragraphs[0].runs[-1].text = "2"
    t.rows[0].cells[1].paragraphs[0].runs[0].text = TITLE
    for r in t.rows[0].cells[1].paragraphs[0].runs[1:]:
        r.text = ""
    dcell = t.rows[1].cells[1]
    if not dcell.paragraphs[0].runs:
        dcell.paragraphs[0].add_run(DATE)
    else:
        dcell.paragraphs[0].runs[0].text = DATE
    gcell = t.rows[2].cells[1]
    if not gcell.paragraphs[0].runs:
        gcell.paragraphs[0].add_run(GITHUB_LINK)
    else:
        gcell.paragraphs[0].runs[0].text = GITHUB_LINK

    set_placeholder(doc, "Aim", AIM)
    set_placeholder(doc, "Description", DESC)
    set_placeholder(doc, "Procedure:", PROCEDURE)

    program_anchor = clear_placeholder(doc, "Program")
    add_code_block(doc, (OUT / "index.html").read_text(encoding="utf8"))
    new_block = doc.paragraphs[-1]
    program_anchor._p.addnext(new_block._p)

    output_anchor = clear_placeholder(doc, "Output")
    insert_after = output_anchor
    for fname, caption in [
        ("output_1_top.jpg", "The empty form, showing every input type"),
        ("output_2_bottom.jpg", "The same form after filling it in and using each control"),
    ]:
        cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(caption); cr.italic = True; cr.font.size = Pt(10)
        img = doc.add_paragraph(); img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img.paragraph_format.space_after = Pt(10)
        img.add_run().add_picture(str(OUT / fname), width=Inches(IMG_W))
        insert_after._p.addnext(cap._p)
        cap._p.addnext(img._p)
        insert_after = img

    set_placeholder(doc, "Result", RESULT)

    for p in list(doc.paragraphs):
        if p.text.strip().startswith("[Kindly Note"):
            p._p.getparent().remove(p._p)

    out_path = OUT / f"Ex2_{URK}.docx"
    doc.save(str(out_path))
    set_ex_number(out_path, "Ex. 2")
    print("wrote", out_path)
    return out_path


if __name__ == "__main__":
    build()
