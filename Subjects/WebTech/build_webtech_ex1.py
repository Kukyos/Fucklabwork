"""Build Cleo's WebTech Ex1 record by filling mam's actual template
(materials/Lab Record Format.docx) with the real HTML5 page, the real code,
and real browser screenshots of it rendered (via claude-in-chrome, not
synthesized) -- same run-for-real philosophy as DSE/DBMS, adapted for a
frontend exercise where "run" means "actually render it in a browser".

    python Subjects/WebTech/build_webtech_ex1.py

index.html + banner.png + the output_*.jpg screenshots must already exist in
output/Ex1/ (written by hand + captured via the browser, not by this script)
-- this script only assembles the docx around them.
"""
from __future__ import annotations

import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from assemble import add_code_block

HERE = Path(__file__).parent
MATERIALS = HERE / "materials"
OUT = HERE / "output" / "Ex1"
URK = "URK24CS1021"
COURSE = "23CS2049 - Web Technology Lab"
TITLE = "Online Gaming Community Website"
DATE = "27/07/2026"
GITHUB_LINK = "https://kukyos.github.io/WebTechLab/Ex1/index.html"
IMG_W = 6.2

AIM = ("To design and develop a structured HTML5 website for an online gaming "
       "community, using semantic elements and formatting techniques to "
       "demonstrate a clear understanding of HTML5 fundamentals.")

DESC = ("The page uses the standard HTML5 structure (doctype, head with "
        "meta/title, body split into sections using <div>). It covers headings "
        "(h1-h3), multiple paragraphs, all six required link types (external "
        "with target=\"_blank\", internal id-based navigation, a bookmark "
        "back-to-top link, a mailto: link, a tel: link, and a clickable image "
        "link), a descriptive <img>, a data table, a description list "
        "(dl/dt/dd), a nested ordered list and a nested unordered list, and an "
        "HTML comment above every section. The background and text colors are "
        "applied with inline CSS on the body and on individual headings, links "
        "and table cells; no external or internal stylesheet is used.")

PROCEDURE = (
    "1. Picked the theme: an online gaming community site (\"Respawn Point\"), "
    "similar in spirit to a Discord community page.\n"
    "2. Created index.html with the <!DOCTYPE html> declaration and the "
    "html/head/meta/title/body skeleton.\n"
    "3. Built the nav bar as a <div> with internal links (#about, #games, #roles, #join, "
    "#rules, #contact) to each section further down the page.\n"
    "4. Added the hero section with the h1 site name and a tagline paragraph.\n"
    "5. Wrote the About section (h2 + h3 + paragraphs) and generated a banner "
    "image, wrapped it in an <a> tag so it's a clickable image link too.\n"
    "6. Added a table listing the featured games, a description list for the "
    "community roles, a nested ordered list for the join steps, and a nested "
    "unordered list for the rules.\n"
    "7. Added the contact section with a mailto: link, a tel: link, an "
    "external Wikipedia link (target=\"_blank\"), and a bookmark link back to "
    "the top of the page.\n"
    "8. Wrote an HTML comment above every section explaining what it is.\n"
    "9. Set the page background color and the text color with inline CSS on "
    "the <body> tag, and used the same inline style attribute on the headings, "
    "links and table header cells to give them the orange and green accent "
    "colors.\n"
    "10. Opened the page in the browser to check every section renders and "
    "every link/id target actually works, then took the output screenshots."
)

RESULT = ("The HTML5 web page for the online gaming community theme was "
          "successfully designed and executed using the required tags, "
          "formatting elements and structure. The output matched the "
          "expected layout and content.")


def set_placeholder(doc: Document, marker: str, text: str) -> None:
    """Find the bracketed placeholder paragraph (its text starts with '[')
    that follows the bold label paragraph == marker, and replace its text."""
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.text.strip() == marker:
            ph = paras[i + 1]
            for r in list(ph.runs):
                r.text = ""
            if ph.runs:
                ph.runs[0].text = text
            else:
                ph.add_run(text)
            return
    raise ValueError(f"placeholder for {marker!r} not found")


def clear_placeholder(doc: Document, marker: str) -> "docx.text.paragraph.Paragraph":
    """Same lookup as set_placeholder, but empties the paragraph (for
    Program/Output where real content gets inserted right after it instead)
    and returns it as the anchor to insert after."""
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.text.strip() == marker:
            ph = paras[i + 1]
            for r in list(ph.runs):
                r.text = ""
            return ph
    raise ValueError(f"placeholder for {marker!r} not found")


def fill_header_regno(doc: Document, urk: str) -> None:
    """The header's course-title cell is a Word content control bound to the
    document's core-properties Title -- doc.sections[0].header.tables misses
    it entirely (python-docx's Table.rows/cells only finds <w:tc> that are
    direct children of <w:tr>; this one is nested inside <w:sdt><w:sdtContent>).
    The [regno] cell next to it is a plain <w:tc> though, so just walk the
    header XML directly for the literal '[regno]' run instead of going
    through the table API."""
    from docx.oxml.ns import qn
    header = doc.sections[0].header._element
    for t in header.iter(qn("w:t")):
        if t.text == "[regno]":
            t.text = urk
            return
    raise ValueError("[regno] placeholder not found in header")


def fill_footer_title(doc: Document, ex_label: str, title: str) -> None:
    """Footer is 'Ex. N | Title of exercise' | <PAGE field>, in a 2-column
    table. 'Ex. N' is itself a content control bound to the extended
    'Company' property (set for real post-save via set_ex_number, since
    that's what Word actually renders) but its CACHED run text also needs
    updating to match, in case anything reads the cached text instead of
    resolving the binding (e.g. some PDF converters). Both it and the plain
    ' | Title of exercise' run are found by walking w:t directly, not the
    table/run API -- same reason as fill_header_regno."""
    from docx.oxml.ns import qn
    footer = doc.sections[0].footer._element
    found_ex, found_title = False, False
    for t in footer.iter(qn("w:t")):
        if t.text == "Ex. 1":
            t.text = ex_label
            found_ex = True
        elif t.text and t.text.strip() == "| Title of exercise":
            t.text = f" | {title}"
            found_title = True
    if not (found_ex and found_title):
        raise ValueError(f"footer placeholders not found (ex={found_ex}, title={found_title})")


def set_ex_number(docx_path: Path, label: str) -> None:
    """The footer's 'Ex. N' text is a content control bound to the document's
    extended 'Company' property (docProps/app.xml), not the cached run text
    -- same class of bug as the header's Title binding. python-docx has no
    high-level API for extended properties, so patch the saved zip directly."""
    import re, zipfile
    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path, "r") as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                text = data.decode("utf8")
                text = re.sub(r"<Company>.*?</Company>", f"<Company>{label}</Company>", text)
                data = text.encode("utf8")
            zout.writestr(item, data)
    tmp.replace(docx_path)


def build() -> Path:
    tmpl = next(MATERIALS.glob("Lab Record Format.docx"))
    doc = Document(str(tmpl))

    # the header's course-title cell is a content control bound to the doc's
    # Title property. The template's original author (Rose, per docProps)
    # left that property as "24CS2502 - MERN Full Stack Development Lab"
    # from whatever course this template was copied from -- Word renders the
    # LIVE bound property, not the cached header text, so it kept showing
    # MERN on every page until this is fixed at the property level
    doc.core_properties.title = COURSE
    doc.core_properties.author = URK
    doc.core_properties.last_modified_by = URK
    fill_header_regno(doc, URK)
    fill_footer_title(doc, "Ex. 1", TITLE)

    # Ex.No/Title/Date/GitHub table
    t = doc.tables[0]
    t.rows[0].cells[1].paragraphs[0].runs[0].text = TITLE
    for r in t.rows[0].cells[1].paragraphs[0].runs[1:]:
        r.text = ""
    dcell = t.rows[1].cells[1]
    if not dcell.paragraphs[0].runs:
        dcell.paragraphs[0].add_run(DATE)
    else:
        dcell.paragraphs[0].runs[0].text = DATE
    gcell = t.rows[2].cells[1]
    if not gcell.paragraphs[0].runs:
        gcell.paragraphs[0].add_run(GITHUB_LINK)
    else:
        gcell.paragraphs[0].runs[0].text = GITHUB_LINK

    set_placeholder(doc, "Aim", AIM)
    set_placeholder(doc, "Description", DESC)
    set_placeholder(doc, "Procedure:", PROCEDURE)

    program_anchor = clear_placeholder(doc, "Program")
    add_code_block(doc, (OUT / "index.html").read_text(encoding="utf8"))
    # add_code_block appends at doc end; move it right after the anchor
    new_block = doc.paragraphs[-1]
    program_anchor._p.addnext(new_block._p)

    output_anchor = clear_placeholder(doc, "Output")
    insert_after = output_anchor
    for fname, caption in [
        ("output_1_top.jpg", "Nav bar, hero, about section, banner image and the games table"),
        ("output_2_bottom.jpg", "Community roles, join steps, ground rules and the contact section"),
    ]:
        cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cr = cap.add_run(caption); cr.italic = True; cr.font.size = Pt(10)
        img = doc.add_paragraph(); img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img.paragraph_format.space_after = Pt(10)
        img.add_run().add_picture(str(OUT / fname), width=Inches(IMG_W))
        insert_after._p.addnext(cap._p)
        cap._p.addnext(img._p)
        insert_after = img

    set_placeholder(doc, "Result", RESULT)

    # drop the "Kindly Note" instruction line, it was for whoever fills the
    # template by hand
    for p in list(doc.paragraphs):
        if p.text.strip().startswith("[Kindly Note"):
            p._p.getparent().remove(p._p)

    out_path = OUT / f"Ex1_{URK}.docx"
    doc.save(str(out_path))
    set_ex_number(out_path, "Ex. 1")
    print("wrote", out_path)
    return out_path


if __name__ == "__main__":
    build()
