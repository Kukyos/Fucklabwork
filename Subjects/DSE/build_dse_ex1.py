"""Build Cleo's DSE Ex-1 record (NumPy & Pandas) from scratch per the
Record Guidelines: run the real Python, capture real output, render terminal
screenshots, assemble the .docx exactly in the guideline's layout.

    python Subjects/DSE/build_dse_ex1.py
"""
from __future__ import annotations

import ast, contextlib, io, json, subprocess, sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from assemble import add_code_block
import labshot

HERE = Path(__file__).parent
OUT = HERE / "output" / "Ex1"
URK = "URK24CS1021"
DATE = "23/07/2026"
COURSE = "23CS2025 DATA SCIENCE ECOSYSTEM LAB"
TITLE = "WORKING WITH NUMPY AND PANDAS"
FONT = "Times New Roman"

AIM = ("To develop and implement Python programs that solves real-world problems in "
       "various industries such as HR, logistics, retail, finance, and security using "
       "Numpy and Pandas.")
DESC = ("NumPy (Numerical Python) is a library for efficient numerical computing. It "
        "provides the ndarray, an n-dimensional array object, together with vectorized "
        "operations for indexing, slicing, reshaping and mathematical computation that "
        "are much faster than native Python lists. Pandas is built on top of NumPy and "
        "provides two main data structures \u2014 the Series (a one-dimensional labelled "
        "array) and the DataFrame (a two-dimensional labelled table) \u2014 for handling "
        "and analysing structured, tabular data.")
QUESTION = (
    "Manufacturing Quality Control\n"
    "Analyze hourly defect rates (%) for 24-hour production:\n"
    "\u2022 Create array of defects (random 0.1-5.0)\n"
    "\u2022 Get rate at hour 12 (indexing)\n"
    "\u2022 Slice night shift (20-6)\n"
    "\u2022 Reshape into (3,8) for shift analysis\n"
    "\u2022 Iterate to find problem hours (>3%)\n"
    "\u2022 Join with temp/humidity array (random 24 values)\n"
    "\u2022 Split into 3 shifts\n"
    "\u2022 Find critical defects (>4%)\n"
    "\u2022 Sort defects in ascending order\n"
    "\u2022 Filter acceptable range (0.5-1.5%)\n"
    "Pandas:\n"
    "\u2022 Convert list of products [\"Widget A\", \"Widget B\", \"Gadget\"] to Series\n"
    "\u2022 Convert {\"Product\": [\"Tool X\", \"Tool Y\"], \"Tolerance\": [0.1, 0.2]} to DataFrame")
RESULT = "The exercise shown above has been successfully executed and the output has been verified."


def run(pyfile: str) -> str:
    r = subprocess.run([sys.executable, str(OUT / pyfile)], capture_output=True, text=True)
    return (r.stdout + r.stderr).rstrip("\n")


def shot(pyfile: str, output: str) -> bytes:
    return labshot.render_terminal_shot(
        output, command=f"python {pyfile}", cwd=r"C:\DSE\Ex1", style="cmd",
        urk_mode="none", show_banner=False, trailing_prompt=True, width=1500)


# ---- notebook + video script ---------------------------------------------
# (kind, source, spoken-note). The spoken note per code cell drives the video
# script, so the notebook and the talk stay in lock-step.
NB_CELLS = [
    ("md", "# Ex 1 — Working with NumPy and Pandas\n\n"
           "**Scenario:** Manufacturing Quality Control  \n**Reg No:** URK24CS1021", None),
    ("md", "### NumPy", None),
    ("code",
     "import numpy as np\nnp.random.seed(1021)\n\n"
     "# hourly defect rates for a full day (24 values, 0.1 - 5.0 %)\n"
     "defects = np.round(np.random.uniform(0.1, 5.0, 24), 2)\ndefects",
     "First I import numpy and set a seed, so the random numbers come out the "
     "same every time I run it. Then I make an array of 24 defect rates — one "
     "for each hour of the day — somewhere between 0.1 and 5 percent, and round "
     "them to two decimals so it's easy to read."),
    ("code", "# defect rate at the 12th hour\ndefects[12]",
     "Here I just index into the array with square brackets to pull out the "
     "defect rate at hour 12."),
    ("code",
     "# night shift runs 20:00 to 06:00, so it wraps past midnight\n"
     "night = np.concatenate((defects[20:], defects[:7]))\nnight",
     "The night shift runs from 8 in the evening to 6 in the morning, which "
     "wraps around midnight. So I slice the last few hours of the day and the "
     "first few hours of the next, and concatenate them into one array."),
    ("code", "# 3 shifts of 8 hours each\nshifts = defects.reshape(3, 8)\nshifts",
     "Next I reshape the 24 hours into a 3 by 8 grid — three shifts of eight "
     "hours each — so I can analyse it shift by shift instead of one long line."),
    ("code",
     "# hours where defects went above 3%\n"
     "for hour, rate in enumerate(defects):\n"
     "    if rate > 3:\n        print(\"hour\", hour, \"->\", rate, \"%\")",
     "Then I loop through every hour with enumerate, and print out only the ones "
     "where the defect rate went above 3 percent — those are my problem hours "
     "that need attention."),
    ("code",
     "# pair each hour's defects with a temperature reading\n"
     "temp = np.round(np.random.uniform(18, 30, 24), 1)\n"
     "combined = np.vstack((defects, temp))\ncombined",
     "Here I make a second array of temperatures and stack it on top of the "
     "defects with vstack, so now each hour has both its defect rate and its "
     "temperature in one 2 by 24 array."),
    ("code",
     "s1, s2, s3 = np.split(defects, 3)\n"
     "print(\"shift 1:\", s1)\nprint(\"shift 2:\", s2)\nprint(\"shift 3:\", s3)",
     "After that I split the day into three equal shifts and print each one out "
     "on its own line."),
    ("code", "# anything above 4% counts as critical\ndefects[defects > 4]",
     "Using a boolean mask, I keep only the critical defects — the hours where "
     "the rate went above 4 percent."),
    ("code", "np.sort(defects)",
     "This one line sorts all the defect rates in ascending order, smallest to "
     "largest."),
    ("code",
     "# acceptable quality is between 0.5 and 1.5%\n"
     "defects[(defects >= 0.5) & (defects <= 1.5)]",
     "And here I filter for the hours in the acceptable range — between 0.5 and "
     "1.5 percent — by combining two conditions with an and."),
    ("md", "### Pandas", None),
    ("code",
     "import pandas as pd\n\n"
     "products = [\"Widget A\", \"Widget B\", \"Gadget\"]\npd.Series(products)",
     "Now switching to pandas. I take a plain list of product names and turn it "
     "into a Series, which just gives each product a little index next to it."),
    ("code",
     "data = {\"Product\": [\"Tool X\", \"Tool Y\"], \"Tolerance\": [0.1, 0.2]}\n"
     "df = pd.DataFrame(data)\ndf",
     "And finally I build a DataFrame from a dictionary, so the products and "
     "their tolerance values line up neatly in a table with rows and columns."),
]


def _srclist(text: str) -> list[str]:
    lines = text.splitlines(keepends=True)
    return lines or [""]


def _run_cell(src: str, ns: dict, count: int) -> list[dict]:
    """Exec a code cell in shared ns like a Jupyter kernel: capture stdout and,
    if the cell ends in a bare expression, its value as an execute_result."""
    tree = ast.parse(src)
    last = tree.body[-1] if tree.body else None
    buf, result = io.StringIO(), None
    with contextlib.redirect_stdout(buf):
        if isinstance(last, ast.Expr):
            exec(compile(ast.Module(tree.body[:-1], []), "<cell>", "exec"), ns)
            result = eval(compile(ast.Expression(last.value), "<cell>", "eval"), ns)
        else:
            exec(compile(tree, "<cell>", "exec"), ns)
    outs = []
    if buf.getvalue():
        outs.append({"output_type": "stream", "name": "stdout",
                     "text": _srclist(buf.getvalue())})
    if result is not None:
        data = {"text/plain": _srclist(repr(result))}
        if hasattr(result, "_repr_html_"):
            data["text/html"] = _srclist(result._repr_html_())
        outs.append({"output_type": "execute_result", "execution_count": count,
                     "metadata": {}, "data": data})
    return outs


def build_notebook(path: Path) -> None:
    ns: dict = {}
    cells, count = [], 0
    for kind, src, _note in NB_CELLS:
        if kind == "md":
            cells.append({"cell_type": "markdown", "metadata": {}, "source": _srclist(src)})
        else:
            count += 1
            cells.append({"cell_type": "code", "execution_count": count,
                          "metadata": {}, "outputs": _run_cell(src, ns, count),
                          "source": _srclist(src)})
    nb = {"cells": cells,
          "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python",
                                       "name": "python3"},
                       "language_info": {"name": "python", "version": "3.14"}},
          "nbformat": 4, "nbformat_minor": 5}
    path.write_text(json.dumps(nb, indent=1), encoding="utf8")


def write_video_script(path: Path) -> None:
    notes = [n for k, _s, n in NB_CELLS if k == "code" and n]
    parts = [
        "# Ex 1 — Code Explanation (video script)",
        "*Manufacturing Quality Control · URK24CS1021 — aim ~5 minutes, spoken casually.*",
        "",
        "**Intro**",
        "Hi, I'm [name], register number URK24CS1021. This is Experiment 1 of the "
        "Data Science Ecosystem Lab, working with NumPy and Pandas. My scenario is "
        "Manufacturing Quality Control — I'm analysing the hourly defect rates from "
        "a factory over a 24-hour production day. Let me walk through my notebook "
        "cell by cell.",
        "",
    ]
    for i, note in enumerate(notes, 1):
        parts += [f"**Cell {i}**", note, ""]
    parts += [
        "**Outro**",
        "So that's the whole notebook — I started with raw hourly defect data and "
        "used NumPy to index, slice, reshape, filter and sort it, then used Pandas "
        "to organise product information into a Series and a DataFrame. Everything "
        "ran and the output was verified. Thanks for watching.",
    ]
    path.write_text("\n".join(parts), encoding="utf8")


# ---- docx helpers ---------------------------------------------------------
def rfonts(el, name=FONT):
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)


def cell_border(cell):
    pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{s}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    pr.append(b)


def label(doc, text, size=12, space_before=8):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.name = FONT; r.font.size = Pt(size)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(3)
    return p


def body(doc, text, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(12); r.italic = italic
    p.paragraph_format.space_after = Pt(4)
    return p


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "screenshots").mkdir(exist_ok=True)
    np_out, pd_out = run("qc_numpy.py"), run("qc_pandas.py")
    np_png, pd_png = shot("qc_numpy.py", np_out), shot("qc_pandas.py", pd_out)
    (OUT / "screenshots" / "numpy_output.png").write_bytes(np_png)
    (OUT / "screenshots" / "pandas_output.png").write_bytes(pd_png)

    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = FONT; normal.font.size = Pt(12)
    rfonts(normal.element)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1.0))

    # header: course (left) + reg no (right), bold
    h = sec.header.paragraphs[0]
    h.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
    r = h.add_run(COURSE); r.bold = True
    h.add_run("\t"); r2 = h.add_run(URK); r2.bold = True
    pbdr = h._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "12")
    bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), "000000")
    bd.append(bot); pbdr.append(bd)
    # footer
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = f.add_run(f"EX-1: {TITLE}"); fr.bold = True

    # Ex.No / Title table
    t = doc.add_table(rows=2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    rows = [("Ex. No. 1", TITLE), ("Date of Exercise", DATE)]
    widths = (Inches(1.9), Inches(4.35))
    for ri, (a, b) in enumerate(rows):
        for ci, txt in enumerate((a, b)):
            c = t.rows[ri].cells[ci]; c.width = widths[ci]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = p.add_run(txt); cr.font.name = FONT; cr.font.size = Pt(12)
            cr.bold = (ci == 0 or ri == 0)
            cell_border(c)
    doc.add_paragraph()

    label(doc, "Aim"); body(doc, AIM)
    label(doc, "Description (About Numpy & Pandas)"); body(doc, DESC)
    label(doc, "Question:"); body(doc, QUESTION)

    label(doc, "Sample Code:")
    add_code_block(doc, (OUT / "qc_numpy.py").read_text())
    label(doc, "Sample Output:")
    _img(doc, np_png)

    label(doc, "Sample Code:")
    add_code_block(doc, (OUT / "qc_pandas.py").read_text())
    label(doc, "Sample Output:")
    _img(doc, pd_png)

    label(doc, "CODE EXPLANATION LINK:")
    body(doc, "(Youtube / Drive Link)", italic=True)
    label(doc, "Preparation course certificate")
    doc.add_paragraph()  # space for the certificate image
    label(doc, "RESULT"); body(doc, RESULT)

    out_docx = OUT / f"Ex1_{URK}.docx"
    doc.save(str(out_docx))

    build_notebook(OUT / "Ex1_QualityControl.ipynb")
    write_video_script(OUT / "video_script.md")
    print("wrote", OUT)
    print("  screenshots: 2 | docx:", out_docx.name,
          "| notebook + video_script.md")


def _img(doc, png, width_in=6.2):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(io.BytesIO(png), width=Inches(width_in))


if __name__ == "__main__":
    build()
