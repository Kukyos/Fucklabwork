"""Shared record builder for the DSE experiments.

build_dse_ex1..4.py each carried their own byte-identical copy of the docx
helpers, the notebook kernel and the docx skeleton. From Ex-5 onward the
per-experiment script only declares its content (title, aim, question,
SECTIONS, NB_CELLS, video script) and calls `build(Exp(...))`.

ex1-4 are left alone - they already ran and their output is checked in.
"""
from __future__ import annotations

import ast, base64, contextlib, io, json, os, sys
from dataclasses import dataclass, field
from pathlib import Path

import matplotlib
matplotlib.use("Agg")  # headless: this runs from a build script, never a GUI
import matplotlib.pyplot as plt

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from assemble import add_code_block
import labshot

URK = "URK24CS1021"
COURSE = "23CS2025 DATA SCIENCE ECOSYSTEM LAB"
FONT = "Times New Roman"
RESULT = "The exercise shown above has been successfully executed and the output has been verified."


@dataclass
class Exp:
    """Everything that differs between one experiment and the next."""
    num: int
    title: str           # e.g. "REGRESSION ANALYSIS"
    date: str            # dd/mm/yyyy
    out: Path            # output/ExN
    aim: str
    desc_label: str      # "Description (About Regression)"
    desc: str
    question: str
    preamble: str        # imports, run once before the sections
    sections: list[tuple[str, str]]      # (heading, code)
    nb_cells: list[tuple[str, str]]      # (kind, source)
    video: str
    script_name: str     # the .py written next to the record
    notebook_name: str
    shot_command: str    # what the fake terminal claims to be running
    shot_cwd: str
    extra_files: dict = field(default_factory=dict)  # name -> bytes/str written into out/


# ---- running the real code ------------------------------------------------
def _grab_figure() -> bytes | None:
    """Save whatever figure was just drawn, then close it. Figure-level seaborn
    calls (pairplot) make their own figure, which is the one gcf() returns."""
    if not plt.get_fignums():
        return None
    buf = io.BytesIO()
    plt.gcf().savefig(buf, format="png", dpi=130, bbox_inches="tight")
    plt.close("all")
    return buf.getvalue()


def run_sections(exp: Exp) -> list[list[bytes]]:
    """Exec each section cumulatively in one shared namespace, so state carries
    forward like a real script. Each section gives back its images: a terminal
    shot if it printed anything, plus its plot if it drew one."""
    ns: dict = {}
    prev_cwd = os.getcwd()
    os.chdir(exp.out)  # the script reads its csv from next to itself
    try:
        exec(compile(exp.preamble, "<preamble>", "exec"), ns)
        shots = []
        for heading, code in exp.sections:
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                exec(compile(code, "<section>", "exec"), ns)
            imgs = []
            if buf.getvalue().strip():
                imgs.append(labshot.render_terminal_shot(
                    buf.getvalue().rstrip("\n"), command=exp.shot_command,
                    cwd=exp.shot_cwd, style="cmd", urk_mode="none",
                    show_banner=False, trailing_prompt=True, width=1500))
            fig = _grab_figure()
            if fig:
                imgs.append(fig)
            assert imgs, f"section produced no output: {heading}"
            shots.append(imgs)
    finally:
        os.chdir(prev_cwd)
    return shots


# ---- notebook -------------------------------------------------------------
def _srclist(text: str) -> list[str]:
    return text.splitlines(keepends=True) or [""]


def _run_cell(src: str, ns: dict, count: int) -> list[dict]:
    tree = ast.parse(src)
    last = tree.body[-1] if tree.body else None
    buf, result = io.StringIO(), None
    with contextlib.redirect_stdout(buf):
        if isinstance(last, ast.Expr):
            exec(compile(ast.Module(tree.body[:-1], []), "<cell>", "exec"), ns)
            result = eval(compile(ast.Expression(last.value), "<cell>", "eval"), ns)
        else:
            exec(compile(tree, "<cell>", "exec"), ns)
    outs = []
    if buf.getvalue():
        outs.append({"output_type": "stream", "name": "stdout",
                     "text": _srclist(buf.getvalue())})
    fig = _grab_figure()  # plot cells return None from plt.show(), so catch the figure
    if fig:
        outs.append({"output_type": "display_data", "metadata": {},
                     "data": {"image/png": base64.b64encode(fig).decode()}})
    if result is not None:
        data = {"text/plain": _srclist(repr(result))}
        if hasattr(result, "_repr_html_"):
            data["text/html"] = _srclist(result._repr_html_())
        outs.append({"output_type": "execute_result", "execution_count": count,
                     "metadata": {}, "data": data})
    return outs


def build_notebook(exp: Exp, path: Path) -> None:
    ns: dict = {}
    cells, count = [], 0
    for kind, src in exp.nb_cells:
        if kind == "md":
            cells.append({"cell_type": "markdown", "metadata": {}, "source": _srclist(src)})
        else:
            count += 1
            # teacher's rule from Ex-2 onward: reg no printed first, every cell
            src = f'print("{URK}")\n\n' + src
            cells.append({"cell_type": "code", "execution_count": count,
                          "metadata": {}, "outputs": _run_cell(src, ns, count),
                          "source": _srclist(src)})
    nb = {"cells": cells,
          "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python",
                                      "name": "python3"},
                       "language_info": {"name": "python", "version": "3.14"}},
          "nbformat": 4, "nbformat_minor": 5}
    path.write_text(json.dumps(nb, indent=1), encoding="utf8")


# ---- docx helpers ---------------------------------------------------------
def rfonts(el, name=FONT):
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)


def cell_border(cell):
    pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{s}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    pr.append(b)


def label(doc, text, size=12, space_before=8):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.name = FONT; r.font.size = Pt(size)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(3)
    return p


def body(doc, text, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(12); r.italic = italic
    p.paragraph_format.space_after = Pt(4)
    return p


def _img(doc, png, width_in=6.2):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(io.BytesIO(png), width=Inches(width_in))


# ---- the whole build ------------------------------------------------------
def build(exp: Exp) -> Path:
    exp.out.mkdir(parents=True, exist_ok=True)
    (exp.out / "screenshots").mkdir(exist_ok=True)
    for name, content in exp.extra_files.items():
        p = exp.out / name
        p.write_bytes(content) if isinstance(content, bytes) else \
            p.write_text(content, encoding="utf8")

    full_code = exp.preamble + "\n\n".join(
        f"# {heading}\n{code}" for heading, code in exp.sections)
    (exp.out / exp.script_name).write_text(
        f"# Ex {exp.num} - {exp.title.title()}\n# {URK}\n\n" + full_code, encoding="utf8")

    shots = run_sections(exp)
    for i, imgs in enumerate(shots, 1):
        for j, png in enumerate(imgs):
            suffix = "" if len(imgs) == 1 else chr(ord("a") + j)
            (exp.out / "screenshots" / f"q{i:02d}{suffix}_output.png").write_bytes(png)

    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = FONT; normal.font.size = Pt(12)
    rfonts(normal.element)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1.0))

    h = sec.header.paragraphs[0]
    h.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
    r = h.add_run(COURSE); r.bold = True
    h.add_run("\t"); r2 = h.add_run(URK); r2.bold = True
    pbdr = h._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "12")
    bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), "000000")
    bd.append(bot); pbdr.append(bd)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = f.add_run(f"EX-{exp.num}: {exp.title}"); fr.bold = True

    t = doc.add_table(rows=2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    rows = [(f"Ex. No. {exp.num}", exp.title), ("Date of Exercise", exp.date)]
    widths = (Inches(1.9), Inches(4.35))
    for ri, (a, b) in enumerate(rows):
        for ci, txt in enumerate((a, b)):
            c = t.rows[ri].cells[ci]; c.width = widths[ci]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = p.add_run(txt); cr.font.name = FONT; cr.font.size = Pt(12)
            cr.bold = (ci == 0 or ri == 0)
            cell_border(c)
    doc.add_paragraph()

    label(doc, "Aim"); body(doc, exp.aim)
    label(doc, exp.desc_label); body(doc, exp.desc)
    label(doc, "Question:"); body(doc, exp.question)

    for (heading, code), imgs in zip(exp.sections, shots):
        label(doc, f"Sample Code: {heading}")
        add_code_block(doc, code)
        label(doc, "Sample Output:")
        for png in imgs:
            _img(doc, png)

    label(doc, "CODE EXPLANATION LINK:")
    body(doc, "(Youtube / Drive Link)", italic=True)
    label(doc, "Preparation course certificate")
    doc.add_paragraph()
    label(doc, "RESULT"); body(doc, RESULT)

    out_docx = exp.out / f"Ex{exp.num}_{URK}.docx"
    doc.save(str(out_docx))

    prev_cwd = os.getcwd()
    os.chdir(exp.out)
    try:
        build_notebook(exp, exp.out / exp.notebook_name)
    finally:
        os.chdir(prev_cwd)
    (exp.out / "video_script.md").write_text(exp.video, encoding="utf8")
    print("wrote", exp.out)
    print("  sections:", len(shots), "| images:", sum(len(i) for i in shots),
          "| docx:", out_docx.name, "|", exp.notebook_name, "+ video_script.md")
    return out_docx
