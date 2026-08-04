"""Build Cleo's DSE Ex-2 record (Pandas / fan speed dataset) from scratch,
same approach as build_dse_ex1.py: run the real Python, capture real output,
render terminal screenshots, assemble the .docx.

    python Subjects/DSE/build_dse_ex2.py
"""
from __future__ import annotations

import ast, contextlib, io, json, os, sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from assemble import add_code_block
import labshot

HERE = Path(__file__).parent
OUT = HERE / "output" / "Ex2"
URK = "URK24CS1021"
DATE = "24/07/2026"
COURSE = "23CS2025 DATA SCIENCE ECOSYSTEM LAB"
TITLE = "WORKING WITH PANDAS"
FONT = "Times New Roman"

AIM = ("To develop and implement Python programs that solves real-world problems in "
       "various industries such as HR, logistics, retail, finance, and security using "
       "Numpy and Pandas.")
DESC = ("Pandas is a Python library built on top of NumPy for handling and analysing "
        "structured, tabular data. Its core structure, the DataFrame, is a "
        "two-dimensional labelled table that supports importing data from files, "
        "inspecting and cleaning it, slicing it by rows/columns/conditions, deriving "
        "new columns, sorting, grouping, building pivot tables, and ranking values \u2014 "
        "the standard steps of exploring any real-world dataset.")
QUESTION = (
    "Dataset: fanspeed.csv (https://dataset.karunya.edu/CSE/fanspeed.csv)\n"
    "Columns: Temperature, Observed Fan Speed\n"
    "\u2022 Import the dataset\n"
    "\u2022 Display the first 10 rows / last 8 rows\n"
    "\u2022 Display information about the dataset (info, column names, shape, "
    "statistical inferences, data types)\n"
    "\u2022 Examine the dataset for null values, display the null count per column, "
    "impute the null values with a suitable method\n"
    "\u2022 Slicing: single/multiple column, single/multiple row, loc & iloc "
    "(single & multiple), conditional slicing\n"
    "\u2022 Add a column to the end using a normal arithmetic operation, apply() with "
    "a custom function, apply() with an inbuilt function\n"
    "\u2022 Add a row and delete it\n"
    "\u2022 Sort the dataset in ascending and descending order\n"
    "\u2022 Group the dataset using single and multiple conditions\n"
    "\u2022 Create a pivot table for the dataset\n"
    "\u2022 Create a ranking column using all the ranking methods")
RESULT = "The exercise shown above has been successfully executed and the output has been verified."

PREAMBLE = (
    "import pandas as pd\n\n"
    "pd.set_option(\"display.max_rows\", 15)\n"
    "pd.set_option(\"display.width\", 120)\n"
    "pd.set_option(\"display.max_columns\", 10)\n\n"
)

# (heading for the docx label, source code for that question)
SECTIONS: list[tuple[str, str]] = [
    ("1. Import the dataset",
     'df = pd.read_csv("fanspeed.csv")\nprint(df)'),

    ("2. Display the first 10 rows and last 8 rows",
     'print("First 10 rows:")\nprint(df.head(10))\n\n'
     'print("\\nLast 8 rows:")\nprint(df.tail(8))'),

    ("3. Display information about the dataset",
     'print("Column names:", list(df.columns))\n'
     'print("\\nShape:", df.shape)\n'
     'print("\\nStatistical inferences:")\nprint(df.describe())\n'
     'print("\\nData types:")\nprint(df.dtypes)\n'
     'print("\\nInfo:")\ndf.info()'),

    ("4. Null values: examine, display, impute",
     '# fanspeed.csv has no missing values as downloaded, so plant a few on a\n'
     '# copy (fixed rows, reproducible) just to show detection + imputation\n'
     '# -- the real df used from here on stays untouched\n'
     'demo = df.copy()\n'
     'demo.loc[[3, 27, 61], "Observed Fan Speed"] = None\n'
     'print("Null values per column:")\nprint(demo.isnull().sum())\n\n'
     'mean_speed = demo["Observed Fan Speed"].mean()\n'
     'demo["Observed Fan Speed"] = demo["Observed Fan Speed"].fillna(mean_speed)\n'
     'print("\\nAfter imputing with the column mean:")\nprint(demo.isnull().sum())'),

    ("5. Slicing the dataset",
     'print("Single column:")\nprint(df["Temperature"].head())\n\n'
     'print("\\nMultiple columns:")\nprint(df[["Temperature", "Observed Fan Speed"]].head())\n\n'
     'print("\\nSingle row (loc):")\nprint(df.loc[5])\n\n'
     'print("\\nMultiple rows (loc):")\nprint(df.loc[5:8])\n\n'
     'print("\\nSingle row (iloc):")\nprint(df.iloc[5])\n\n'
     'print("\\nMultiple rows (iloc):")\nprint(df.iloc[5:8])\n\n'
     'print("\\nloc - rows + one column:")\nprint(df.loc[5:8, "Temperature"])\n\n'
     'print("\\niloc - rows + one column:")\nprint(df.iloc[5:8, 0])\n\n'
     'print("\\nConditional slicing (Temperature > 25):")\n'
     'print(df[df["Temperature"] > 25].head())'),

    ("6. Add a column to the end",
     '# normal arithmetic operation\n'
     'df["Speed_Per_Degree"] = df["Observed Fan Speed"] / df["Temperature"]\n'
     '# apply() with a custom function\n'
     'df["Speed_Level"] = df["Observed Fan Speed"].apply(lambda x: "High" if x > 60 else "Low")\n'
     '# apply() with an inbuilt function\n'
     'df["Speed_Rounded"] = df["Observed Fan Speed"].apply(round)\n'
     'print(df.head())'),

    ("7. Add a row and delete it",
     'new_row = pd.DataFrame([{"Temperature": 40, "Observed Fan Speed": 120,\n'
     '                          "Speed_Per_Degree": 3.0, "Speed_Level": "High",\n'
     '                          "Speed_Rounded": 120}])\n'
     'df = pd.concat([df, new_row], ignore_index=True)\n'
     'print("After adding a row:")\nprint(df.tail(3))\n\n'
     'df = df.drop(df.index[-1])\n'
     'print("\\nAfter deleting it:")\nprint(df.tail(3))'),

    ("8. Sort ascending and descending",
     'print("Ascending by Observed Fan Speed:")\n'
     'print(df.sort_values("Observed Fan Speed").head())\n\n'
     'print("\\nDescending by Observed Fan Speed:")\n'
     'print(df.sort_values("Observed Fan Speed", ascending=False).head())'),

    ("9. Group by single and multiple conditions",
     'print("Group by Speed_Level (single condition):")\n'
     'print(df.groupby("Speed_Level")["Observed Fan Speed"].mean())\n\n'
     'df["Temp_Band"] = df["Temperature"].apply(\n'
     '    lambda t: "Low" if t < 15 else ("Mid" if t < 25 else "High"))\n'
     'print("\\nGroup by Speed_Level & Temp_Band (multiple conditions):")\n'
     'print(df.groupby(["Speed_Level", "Temp_Band"])["Observed Fan Speed"].mean())'),

    ("10. Pivot table",
     'print(pd.pivot_table(df, index="Temp_Band", columns="Speed_Level",\n'
     '                      values="Observed Fan Speed", aggfunc="mean"))'),

    ("11. Ranking column, all methods",
     'df["Rank_average"] = df["Observed Fan Speed"].rank(method="average")\n'
     'df["Rank_min"] = df["Observed Fan Speed"].rank(method="min")\n'
     'df["Rank_max"] = df["Observed Fan Speed"].rank(method="max")\n'
     'df["Rank_first"] = df["Observed Fan Speed"].rank(method="first")\n'
     'df["Rank_dense"] = df["Observed Fan Speed"].rank(method="dense")\n'
     'print(df[["Observed Fan Speed", "Rank_average", "Rank_min",\n'
     '          "Rank_max", "Rank_first", "Rank_dense"]].head())'),
]


def shot(output: str) -> bytes:
    return labshot.render_terminal_shot(
        output, command="python fs_pandas.py", cwd=r"C:\DSE\Ex2", style="cmd",
        urk_mode="none", show_banner=False, trailing_prompt=True, width=1500)


def run_sections() -> list[bytes]:
    """Exec each section cumulatively in one shared namespace (like a real
    script run top to bottom), capturing each section's own stdout for its
    own screenshot."""
    ns: dict = {}
    prev_cwd = os.getcwd()
    os.chdir(OUT)  # fs_pandas.py reads "fanspeed.csv" relative to itself
    try:
        exec(compile(PREAMBLE, "<preamble>", "exec"), ns)
        shots = []
        for _heading, code in SECTIONS:
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                exec(compile(code, "<section>", "exec"), ns)
            shots.append(shot(buf.getvalue().rstrip("\n")))
    finally:
        os.chdir(prev_cwd)
    return shots


# ---- notebook + video script ---------------------------------------------
NB_CELLS = [
    ("md", "# Ex 2 — Working with Pandas\n\n"
           "**Dataset:** fanspeed.csv  \n**Reg No:** URK24CS1021", None),
    ("code", "import pandas as pd\n\ndf = pd.read_csv(\"fanspeed.csv\")\ndf",
     "First I import pandas and load the fan speed csv into a DataFrame."),
    ("code", "df.head(10)",
     "head with 10 shows the first ten rows."),
    ("code", "df.tail(8)",
     "and tail shows the last eight."),
    ("code", "df.shape, list(df.columns), df.dtypes",
     "A quick look at the shape, the column names, and their data types."),
    ("code", "df.describe()",
     "describe gives the statistical summary — mean, std, min, max, quartiles."),
    ("code",
     "demo = df.copy()\n"
     "demo.loc[[3, 27, 61], \"Observed Fan Speed\"] = None\n"
     "demo.isnull().sum()",
     "The dataset has no real missing values, so I plant a few on purpose, on "
     "a copy, just to show how you'd spot them with isnull and sum."),
    ("code",
     "demo[\"Observed Fan Speed\"] = demo[\"Observed Fan Speed\"].fillna(demo[\"Observed Fan Speed\"].mean())\n"
     "demo.isnull().sum()",
     "Then I fill those gaps with the column mean, and confirm the null count "
     "drops back to zero — the real df is untouched, this was just on the copy."),
    ("code", "df.loc[5:8, [\"Temperature\"]]",
     "loc lets me slice by label — rows 5 through 8, just the Temperature column."),
    ("code", "df[df[\"Temperature\"] > 25].head()",
     "and a boolean condition inside the brackets gives conditional slicing."),
    ("code",
     "df[\"Speed_Per_Degree\"] = df[\"Observed Fan Speed\"] / df[\"Temperature\"]\n"
     "df[\"Speed_Level\"] = df[\"Observed Fan Speed\"].apply(lambda x: \"High\" if x > 60 else \"Low\")\n"
     "df.head()",
     "I add one column with plain arithmetic, and another with apply and a "
     "small lambda that labels each row High or Low."),
    ("code",
     "df.sort_values(\"Observed Fan Speed\", ascending=False).head()",
     "Sorting by fan speed, descending, to see the fastest readings first."),
    ("code",
     "df.groupby(\"Speed_Level\")[\"Observed Fan Speed\"].mean()",
     "Grouping by the Speed_Level column and averaging within each group."),
    ("code",
     "df[\"Rank_dense\"] = df[\"Observed Fan Speed\"].rank(method=\"dense\")\n"
     "df.sort_values(\"Rank_dense\").head()",
     "And finally a ranking column with the dense method, sorted to check it "
     "lines up with the actual values."),
]

CUES = [
    ["import pandas, load fanspeed.csv into a DataFrame"],
    ["head(10) -> first ten rows"],
    ["tail(8) -> last eight rows"],
    ["shape, columns, dtypes -> quick shape check"],
    ["describe() -> mean/std/min/max/quartiles"],
    ["plant a few nulls on purpose", "isnull().sum() to find them"],
    ["fillna with the column mean", "confirm nulls are gone"],
    ["loc by label -> rows 5-8, Temperature column"],
    ["boolean condition in brackets -> conditional slicing"],
    ["arithmetic column", "apply + lambda -> High/Low label column"],
    ["sort_values descending -> fastest readings on top"],
    ["groupby Speed_Level, mean per group"],
    ["rank with dense method", "sort by it to sanity check"],
]


def _srclist(text: str) -> list[str]:
    return text.splitlines(keepends=True) or [""]


def _run_cell(src: str, ns: dict, count: int) -> list[dict]:
    tree = ast.parse(src)
    last = tree.body[-1] if tree.body else None
    buf, result = io.StringIO(), None
    with contextlib.redirect_stdout(buf):
        if isinstance(last, ast.Expr):
            exec(compile(ast.Module(tree.body[:-1], []), "<cell>", "exec"), ns)
            result = eval(compile(ast.Expression(last.value), "<cell>", "eval"), ns)
        else:
            exec(compile(tree, "<cell>", "exec"), ns)
    outs = []
    if buf.getvalue():
        outs.append({"output_type": "stream", "name": "stdout",
                     "text": _srclist(buf.getvalue())})
    if result is not None:
        data = {"text/plain": _srclist(repr(result))}
        if hasattr(result, "_repr_html_"):
            data["text/html"] = _srclist(result._repr_html_())
        outs.append({"output_type": "execute_result", "execution_count": count,
                     "metadata": {}, "data": data})
    return outs


def build_notebook(path: Path) -> None:
    ns: dict = {}
    cells, count = [], 0
    for kind, src, _note in NB_CELLS:
        if kind == "md":
            cells.append({"cell_type": "markdown", "metadata": {}, "source": _srclist(src)})
        else:
            count += 1
            # teacher's rule from Ex-2 onward: reg no printed first, every cell
            src = f'print("{URK}")\n\n' + src
            cells.append({"cell_type": "code", "execution_count": count,
                          "metadata": {}, "outputs": _run_cell(src, ns, count),
                          "source": _srclist(src)})
    nb = {"cells": cells,
          "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python",
                                       "name": "python3"},
                       "language_info": {"name": "python", "version": "3.14"}},
          "nbformat": 4, "nbformat_minor": 5}
    path.write_text(json.dumps(nb, indent=1), encoding="utf8")


def write_video_script(path: Path) -> None:
    p = ["# Ex 2 — Code explanation (talking cues)",
         "*Pandas on fanspeed.csv · URK24CS1021*",
         "",
         "> Cue cards, ~5 min. **Glance, don't read** — look at the cell on screen "
         "and say it in your own words. Keep it relaxed, like you're showing a friend.",
         "",
         "**Open with:** who you are + reg no URK24CS1021, Experiment 2, Pandas, "
         "scenario = exploring a fan speed dataset (temperature vs observed fan speed).",
         ""]
    for i, cues in enumerate(CUES, 1):
        p.append(f"**Cell {i}** — " + "; ".join(cues))
    p += ["",
          "**Close with:** loaded and inspected the data, found and imputed nulls, "
          "sliced it different ways, derived new columns, sorted/grouped/pivoted it, "
          "and ranked it — the standard pandas exploration toolkit."]
    path.write_text("\n".join(p), encoding="utf8")


# ---- docx helpers (same as build_dse_ex1.py) ------------------------------
def rfonts(el, name=FONT):
    rpr = el.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), name)


def cell_border(cell):
    pr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for s in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{s}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        b.append(e)
    pr.append(b)


def label(doc, text, size=12, space_before=8):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = True; r.font.name = FONT; r.font.size = Pt(size)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(3)
    return p


def body(doc, text, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(12); r.italic = italic
    p.paragraph_format.space_after = Pt(4)
    return p


def _img(doc, png, width_in=6.2):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(io.BytesIO(png), width=Inches(width_in))


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "screenshots").mkdir(exist_ok=True)

    full_code = PREAMBLE + "\n\n".join(
        f"# {heading}\n{code}" for heading, code in SECTIONS)
    (OUT / "fs_pandas.py").write_text(
        "# Ex 2 - Working with Pandas (fanspeed.csv)\n# URK24CS1021\n\n" + full_code,
        encoding="utf8")

    shots = run_sections()
    for i, png in enumerate(shots, 1):
        (OUT / "screenshots" / f"q{i:02d}_output.png").write_bytes(png)

    doc = Document()
    normal = doc.styles["Normal"]; normal.font.name = FONT; normal.font.size = Pt(12)
    rfonts(normal.element)
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(1.0))

    h = sec.header.paragraphs[0]
    h.paragraph_format.tab_stops.add_tab_stop(Cm(15.9), WD_TAB_ALIGNMENT.RIGHT)
    r = h.add_run(COURSE); r.bold = True
    h.add_run("\t"); r2 = h.add_run(URK); r2.bold = True
    pbdr = h._p.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "12")
    bot.set(qn("w:space"), "1"); bot.set(qn("w:color"), "000000")
    bd.append(bot); pbdr.append(bd)
    f = sec.footer.paragraphs[0]; f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = f.add_run(f"EX-2: {TITLE}"); fr.bold = True

    t = doc.add_table(rows=2, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    rows = [("Ex. No. 2", TITLE), ("Date of Exercise", DATE)]
    widths = (Inches(1.9), Inches(4.35))
    for ri, (a, b) in enumerate(rows):
        for ci, txt in enumerate((a, b)):
            c = t.rows[ri].cells[ci]; c.width = widths[ci]
            c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = c.paragraphs[0]
            if ri == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cr = p.add_run(txt); cr.font.name = FONT; cr.font.size = Pt(12)
            cr.bold = (ci == 0 or ri == 0)
            cell_border(c)
    doc.add_paragraph()

    label(doc, "Aim"); body(doc, AIM)
    label(doc, "Description (About Pandas)"); body(doc, DESC)
    label(doc, "Question:"); body(doc, QUESTION)

    for (heading, code), png in zip(SECTIONS, shots):
        label(doc, f"Sample Code: {heading}")
        add_code_block(doc, code)
        label(doc, "Sample Output:")
        _img(doc, png)

    label(doc, "CODE EXPLANATION LINK:")
    body(doc, "(Youtube / Drive Link)", italic=True)
    label(doc, "Preparation course certificate")
    doc.add_paragraph()
    label(doc, "RESULT"); body(doc, RESULT)

    out_docx = OUT / f"Ex2_{URK}.docx"
    doc.save(str(out_docx))

    prev_cwd = os.getcwd()
    os.chdir(OUT)
    try:
        build_notebook(OUT / "Ex2_FanSpeed.ipynb")
    finally:
        os.chdir(prev_cwd)
    write_video_script(OUT / "video_script.md")
    print("wrote", OUT)
    print("  screenshots:", len(shots), "| docx:", out_docx.name,
          "| notebook + video_script.md")


if __name__ == "__main__":
    build()
